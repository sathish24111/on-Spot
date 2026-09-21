import os
import re
import io
import json
import random
from datetime import datetime, date
from decimal import Decimal

from flask import Flask, request, jsonify, send_from_directory, send_file
from flask_cors import CORS
from dotenv import load_dotenv
import pymysql
from pymysql.cursors import DictCursor
from dbutils.pooled_db import PooledDB
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

# Load environment configuration
load_dotenv()

# ==========================================
# CONFIGURATION & CONSTANTS
# ==========================================
import urllib.parse
import ssl

PORT = int(os.getenv('PORT', 3000))
DB_HOST = os.getenv('DB_HOST', 'localhost')
DB_PORT = int(os.getenv('DB_PORT', 3306))
DB_USER = os.getenv('DB_USER', 'root')
DB_PASSWORD = os.getenv('DB_PASSWORD', '')
DB_NAME = os.getenv('DB_NAME', 'event_spot_registration')
DB_CONNECTION_LIMIT = int(os.getenv('DB_CONNECTION_LIMIT', 30))

# Support Full Service URI / DATABASE_URL (e.g. from Aiven MySQL)
DATABASE_URL = os.getenv('DATABASE_URL') or os.getenv('MYSQL_URL') or os.getenv('DB_URI')
if DATABASE_URL:
    try:
        url = urllib.parse.urlparse(DATABASE_URL)
        if url.hostname: DB_HOST = url.hostname
        if url.port: DB_PORT = int(url.port)
        if url.username: DB_USER = url.username
        if url.password: DB_PASSWORD = url.password
        if url.path and url.path.strip('/'): DB_NAME = url.path.strip('/')
        print(f"[CONFIG] Parsed DATABASE_URL: Host={DB_HOST}, Port={DB_PORT}, User={DB_USER}, DB={DB_NAME}")
    except Exception as parse_err:
        print(f"[WARNING] Failed to parse DATABASE_URL: {parse_err}")

# Auto-detect SSL requirement (Aiven Cloud or explicit flag)
DB_SSL = os.getenv('DB_SSL', '').lower() in ('true', '1', 'yes', 'required') or ('aivencloud.com' in DB_HOST.lower())

HACKATHON_THEMES = [
    'Intelligent Systems',
    'Sustainability',
    'Tech for Everyday',
    'FinTech / CommerceTech',
    'Smart Business Management / Institution Automation'
]

EVENTS = [
    # Day 1 Events (10)
    {'id': 'd1_adzap', 'name': 'Adzap', 'day': 'day1', 'category': 'Creative & Marketing', 'min_participants': 5, 'max_participants': 5, 'is_standalone': False},
    {'id': 'd1_sharktank', 'name': 'Business Plan Presentation - Shark Tank', 'day': 'day1', 'category': 'Business & Entrepreneurship', 'min_participants': 3, 'max_participants': 5, 'is_standalone': False},
    {'id': 'd1_quiz', 'name': 'Quiz', 'day': 'day1', 'category': 'Academic & Technical', 'min_participants': 2, 'max_participants': 2, 'is_standalone': False},
    {'id': 'd1_designer', 'name': 'Designer Contest', 'day': 'day1', 'category': 'Design & Arts', 'min_participants': 1, 'max_participants': 1, 'is_standalone': False},
    {'id': 'd1_fashion', 'name': 'Fashion Show', 'day': 'day1', 'category': 'Cultural & Performing Arts', 'min_participants': 5, 'max_participants': 12, 'is_standalone': False},
    {'id': 'd1_nailart', 'name': 'Nail Art', 'day': 'day1', 'category': 'Arts & Craft', 'min_participants': 1, 'max_participants': 2, 'is_standalone': False},
    {'id': 'd1_mehandi', 'name': 'Mehandi', 'day': 'day1', 'category': 'Arts & Craft', 'min_participants': 2, 'max_participants': 2, 'is_standalone': False},
    {'id': 'd1_facepainting', 'name': 'Face Painting', 'day': 'day1', 'category': 'Arts & Craft', 'min_participants': 2, 'max_participants': 2, 'is_standalone': False},
    {'id': 'd1_hackathon', 'name': 'Hackathon', 'day': 'day1', 'category': 'Coding & Hackathon', 'min_participants': 3, 'max_participants': 5, 'is_standalone': False, 'has_themes': True},
    {'id': 'd1_freefire', 'name': 'Free Fire Tournament', 'day': 'day1', 'category': 'Gaming', 'min_participants': 4, 'max_participants': 4, 'is_standalone': True},

    # Day 2 Events (11)
    {'id': 'd2_fixbug', 'name': 'Fix The Bug', 'day': 'day2', 'category': 'Technical & Coding', 'min_participants': 1, 'max_participants': 1, 'is_standalone': False},
    {'id': 'd2_webforge', 'name': 'Webforge', 'day': 'day2', 'category': 'Technical & Web', 'min_participants': 1, 'max_participants': 1, 'is_standalone': False},
    {'id': 'd2_paperspark', 'name': 'Paperspark', 'day': 'day2', 'category': 'Paper Presentation', 'min_participants': 1, 'max_participants': 2, 'is_standalone': False},
    {'id': 'd2_prompting', 'name': 'Prompting The Wars', 'day': 'day2', 'category': 'AI & Tech', 'min_participants': 1, 'max_participants': 1, 'is_standalone': False},
    {'id': 'd2_connection', 'name': 'Connection', 'day': 'day2', 'category': 'Fun & Strategy', 'min_participants': 2, 'max_participants': 2, 'is_standalone': False},
    {'id': 'd2_solodance', 'name': 'Solo Dance', 'day': 'day2', 'category': 'Cultural & Performing Arts', 'min_participants': 1, 'max_participants': 1, 'is_standalone': False},
    {'id': 'd2_groupdance', 'name': 'Group Dance', 'day': 'day2', 'category': 'Cultural & Performing Arts', 'min_participants': 8, 'max_participants': 12, 'is_standalone': False},
    {'id': 'd2_solosinging', 'name': 'Solo Singing', 'day': 'day2', 'category': 'Music & Singing', 'min_participants': 1, 'max_participants': 1, 'is_standalone': False},
    {'id': 'd2_photography', 'name': 'Photography Competition', 'day': 'day2', 'category': 'Media & Photography', 'min_participants': 1, 'max_participants': 1, 'is_standalone': False},
    {'id': 'd2_shortfilm', 'name': 'Short Film Competition', 'day': 'day2', 'category': 'Media & Film', 'min_participants': 3, 'max_participants': 5, 'is_standalone': False},
    {'id': 'd2_reels', 'name': 'Reels Creation', 'day': 'day2', 'category': 'Media & Social', 'min_participants': 3, 'max_participants': 5, 'is_standalone': False}
]

PRICING = {
    'day1': 250,
    'day2': 250,
    'both': 350,
    'DAY1_ONLY': 250,
    'DAY2_ONLY': 250,
    'BOTH_DAYS': 350,
    'FREE_FIRE_STANDALONE': 400
}

LIMITS = {
    'DAY1_MAX_EVENTS': 2,
    'DAY2_MAX_EVENTS': 2,
    'BOTH_DAYS_MAX_EVENTS': 4
}

# ==========================================
# PORTAL & EVENT DESK CREDENTIALS (IN-APP)
# ==========================================
# You can change any username or password directly here, and it takes effect immediately!
default_creds = [
    {'key': 'spot_registration', 'name': 'Spot Registration Desk', 'category': 'PORTAL', 'user': 'spot', 'pass': 'spot123'},
    {'key': 'event_issue', 'name': 'Event Issue Management Desk', 'category': 'PORTAL', 'user': 'issue', 'pass': 'issue1234'},
    {'key': 'day2_gate', 'name': 'Day 2 Gate Search Desk', 'category': 'PORTAL', 'user': 'gate', 'pass': 'gate123'},
    {'key': 'coordinator_portal', 'name': 'Coordinator Portal (Master)', 'category': 'PORTAL', 'user': 'coordinator', 'pass': 'coord123'},
    {'key': 'admin_master', 'name': 'Admin Master Portal', 'category': 'PORTAL', 'user': 'admin', 'pass': 'admin123'},

    {'key': 'd1_adzap', 'name': 'Adzap Desk', 'category': 'EVENT', 'user': 'adzap', 'pass': 'adzap123'},
    {'key': 'd1_sharktank', 'name': 'Shark Tank Desk', 'category': 'EVENT', 'user': 'sharktank', 'pass': 'shark123'},
    {'key': 'd1_quiz', 'name': 'Quiz Desk', 'category': 'EVENT', 'user': 'quiz', 'pass': 'quiz123'},
    {'key': 'd1_designer', 'name': 'Designer Contest Desk', 'category': 'EVENT', 'user': 'designer', 'pass': 'design123'},
    {'key': 'd1_fashion', 'name': 'Fashion Show Desk', 'category': 'EVENT', 'user': 'fashion', 'pass': 'fashion123'},
    {'key': 'd1_nailart', 'name': 'Nail Art Desk', 'category': 'EVENT', 'user': 'nailart', 'pass': 'nail123'},
    {'key': 'd1_mehandi', 'name': 'Mehandi Desk', 'category': 'EVENT', 'user': 'mehandi', 'pass': 'mehandi123'},
    {'key': 'd1_facepainting', 'name': 'Face Painting Desk', 'category': 'EVENT', 'user': 'facepainting', 'pass': 'face123'},
    {'key': 'd1_hackathon', 'name': 'Hackathon Desk', 'category': 'EVENT', 'user': 'hackathon', 'pass': 'hack123'},
    {'key': 'd1_freefire', 'name': 'Free Fire Esports Desk', 'category': 'EVENT', 'user': 'freefire', 'pass': 'ff123'},

    {'key': 'd2_fixthebug', 'name': 'Fix The Bug Desk', 'category': 'EVENT', 'user': 'fixthebug', 'pass': 'bug123'},
    {'key': 'd2_webforge', 'name': 'Webforge Desk', 'category': 'EVENT', 'user': 'webforge', 'pass': 'web123'},
    {'key': 'd2_paperspark', 'name': 'Paperspark Desk', 'category': 'EVENT', 'user': 'paperspark', 'pass': 'paper123'},
    {'key': 'd2_prompting', 'name': 'Prompting The Wars Desk', 'category': 'EVENT', 'user': 'prompting', 'pass': 'prompt123'},
    {'key': 'd2_connection', 'name': 'Connection Desk', 'category': 'EVENT', 'user': 'connection', 'pass': 'connect123'},
    {'key': 'd2_solodance', 'name': 'Solo Dance Desk', 'category': 'EVENT', 'user': 'solodance', 'pass': 'dance123'},
    {'key': 'd2_groupdance', 'name': 'Group Dance Desk', 'category': 'EVENT', 'user': 'groupdance', 'pass': 'group123'},
    {'key': 'd2_solosinging', 'name': 'Solo Singing Desk', 'category': 'EVENT', 'user': 'solosinging', 'pass': 'sing123'},
    {'key': 'd2_photography', 'name': 'Photography Desk', 'category': 'EVENT', 'user': 'photography', 'pass': 'photo123'},
    {'key': 'd2_shortfilm', 'name': 'Short Film Desk', 'category': 'EVENT', 'user': 'shortfilm', 'pass': 'film123'},
    {'key': 'd2_reels', 'name': 'Reels Creation Desk', 'category': 'EVENT', 'user': 'reels', 'pass': 'reels123'}
]

def get_cred_val(cred_dict, *keys, default=''):
    for k in keys:
        if k in cred_dict and cred_dict[k] is not None:
            return cred_dict[k]
    return default

# ==========================================
# FLASK APP SETUP
# ==========================================
app = Flask(__name__, static_folder='public', static_url_path='')
CORS(app)

def json_serial(obj):
    if isinstance(obj, (datetime, date)):
        return obj.isoformat()
    if isinstance(obj, Decimal):
        return float(obj)
    raise TypeError(f"Type {type(obj)} not serializable")

app.json.default = json_serial

# ==========================================
# DATABASE CONNECTION POOL & HELPERS
# ==========================================
db_pool = None

def init_db_pool():
    global db_pool
    print(f"[INFO] Connecting to MySQL at {DB_HOST}:{DB_PORT} as '{DB_USER}' (SSL={DB_SSL})...")
    
    ssl_kwargs = {}
    if DB_SSL:
        try:
            ctx = ssl.create_default_context()
            ctx.check_hostname = False
            ctx.verify_mode = ssl.CERT_NONE
            ssl_kwargs['ssl'] = ctx
        except Exception as ssl_err:
            print(f"[WARNING] SSL Context initialization: {ssl_err}")
            ssl_kwargs['ssl'] = {'check_hostname': False}

    # 1. Create Database if needed (safe for cloud databases)
    try:
        init_conn = pymysql.connect(
            host=DB_HOST,
            port=DB_PORT,
            user=DB_USER,
            password=DB_PASSWORD,
            charset='utf8mb4',
            **ssl_kwargs
        )
        with init_conn.cursor() as cur:
            cur.execute(f"CREATE DATABASE IF NOT EXISTS `{DB_NAME}` CHARACTER SET utf8mb4 COLLATE utf8mb4_unicode_ci")
        init_conn.close()
    except Exception as create_err:
        print(f"[INFO] Database existence check: '{DB_NAME}' ({create_err})")

    # 2. Connection Pool
    db_pool = PooledDB(
        creator=pymysql,
        maxconnections=DB_CONNECTION_LIMIT,
        mincached=2,
        maxcached=10,
        blocking=True,
        host=DB_HOST,
        port=DB_PORT,
        user=DB_USER,
        password=DB_PASSWORD,
        database=DB_NAME,
        charset='utf8mb4',
        cursorclass=DictCursor,
        autocommit=True,
        **ssl_kwargs
    )
    print(f"[SUCCESS] Connected to MySQL Database: '{DB_NAME}' (Pool Size: {DB_CONNECTION_LIMIT}, SSL={DB_SSL})")
    init_schema()

def get_connection():
    global db_pool
    if not db_pool:
        init_db_pool()
    return db_pool.connection()

def normalize_sql(sql):
    s = sql.strip()
    s = re.sub(r'INSERT\s+OR\s+REPLACE\s+INTO', 'REPLACE INTO', s, flags=re.IGNORECASE)
    s = re.sub(r'INSERT\s+OR\s+IGNORE\s+INTO', 'INSERT IGNORE INTO', s, flags=re.IGNORECASE)
    return s.replace('?', '%s')

def query_run(sql, params=None):
    norm_sql = normalize_sql(sql)
    conn = get_connection()
    try:
        with conn.cursor() as cur:
            cur.execute(norm_sql, params or ())
            last_id = cur.lastrowid
            affected = cur.rowcount
            return {'lastID': last_id, 'insertId': last_id, 'changes': affected, 'affectedRows': affected}
    finally:
        conn.close()

def query_get(sql, params=None):
    norm_sql = normalize_sql(sql)
    conn = get_connection()
    try:
        with conn.cursor() as cur:
            cur.execute(norm_sql, params or ())
            return cur.fetchone()
    finally:
        conn.close()

def query_all(sql, params=None):
    norm_sql = normalize_sql(sql)
    conn = get_connection()
    try:
        with conn.cursor() as cur:
            cur.execute(norm_sql, params or ())
            return cur.fetchall() or []
    finally:
        conn.close()

def init_schema():
    # 1. events
    query_run("""
        CREATE TABLE IF NOT EXISTS events (
          id VARCHAR(50) PRIMARY KEY,
          name VARCHAR(255) NOT NULL,
          day VARCHAR(20) NOT NULL,
          category VARCHAR(100) NOT NULL,
          is_standalone TINYINT(1) DEFAULT 0,
          min_participants INT DEFAULT 1,
          max_participants INT DEFAULT 1
        ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COLLATE=utf8mb4_unicode_ci
    """)

    # 2. students
    query_run("""
        CREATE TABLE IF NOT EXISTS students (
          id INT AUTO_INCREMENT PRIMARY KEY,
          reg_code VARCHAR(100) UNIQUE NOT NULL,
          name VARCHAR(255) NOT NULL,
          phone VARCHAR(50) NOT NULL,
          email VARCHAR(255) DEFAULT '',
          college VARCHAR(255) NOT NULL,
          day_selection VARCHAR(50) NOT NULL,
          total_fee INT NOT NULL,
          registration_type VARCHAR(50) DEFAULT 'ONLINE',
          food_given TINYINT(1) DEFAULT 0,
          tag_given TINYINT(1) DEFAULT 0,
          food_d2_given TINYINT(1) DEFAULT 0,
          tag_d2_given TINYINT(1) DEFAULT 0,
          counter_name VARCHAR(100) DEFAULT NULL,
          faculty_name VARCHAR(255) DEFAULT NULL,
          registered_at DATETIME DEFAULT CURRENT_TIMESTAMP,
          INDEX idx_students_reg_type (registration_type),
          INDEX idx_students_phone (phone),
          INDEX idx_students_name (name),
          INDEX idx_students_reg_code (reg_code),
          INDEX idx_students_day_sel (day_selection),
          INDEX idx_students_counter (counter_name)
        ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COLLATE=utf8mb4_unicode_ci
    """)

    # 3. student_events
    query_run("""
        CREATE TABLE IF NOT EXISTS student_events (
          id INT AUTO_INCREMENT PRIMARY KEY,
          student_id INT NOT NULL,
          event_id VARCHAR(50) NOT NULL,
          event_day VARCHAR(20) NOT NULL,
          verified_status TINYINT(1) DEFAULT 0,
          event_status VARCHAR(50) DEFAULT 'NOT_STARTED',
          completed_at DATETIME DEFAULT NULL,
          hackathon_theme VARCHAR(255) DEFAULT NULL,
          INDEX idx_se_event_id (event_id),
          INDEX idx_se_student_id (student_id),
          CONSTRAINT fk_se_student FOREIGN KEY (student_id) REFERENCES students(id) ON DELETE CASCADE,
          CONSTRAINT fk_se_event FOREIGN KEY (event_id) REFERENCES events(id) ON DELETE CASCADE
        ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COLLATE=utf8mb4_unicode_ci
    """)

    # 4. freefire_teams
    query_run("""
        CREATE TABLE IF NOT EXISTS freefire_teams (
          id INT AUTO_INCREMENT PRIMARY KEY,
          team_code VARCHAR(100) UNIQUE NOT NULL,
          team_name VARCHAR(255) NOT NULL,
          registration_source VARCHAR(100) NOT NULL,
          counter_name VARCHAR(100) DEFAULT NULL,
          faculty_name VARCHAR(255) DEFAULT NULL,
          registered_at DATETIME DEFAULT CURRENT_TIMESTAMP
        ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COLLATE=utf8mb4_unicode_ci
    """)

    # 5. freefire_players
    query_run("""
        CREATE TABLE IF NOT EXISTS freefire_players (
          id INT AUTO_INCREMENT PRIMARY KEY,
          team_id INT DEFAULT NULL,
          name VARCHAR(255) NOT NULL,
          phone VARCHAR(50) NOT NULL,
          email VARCHAR(255) DEFAULT '',
          college VARCHAR(255) NOT NULL,
          registration_type VARCHAR(50) NOT NULL,
          is_captain TINYINT(1) DEFAULT 0,
          counter_name VARCHAR(100) DEFAULT NULL,
          faculty_name VARCHAR(255) DEFAULT NULL,
          registered_at DATETIME DEFAULT CURRENT_TIMESTAMP,
          INDEX idx_ff_team_id (team_id),
          CONSTRAINT fk_ff_team FOREIGN KEY (team_id) REFERENCES freefire_teams(id) ON DELETE SET NULL
        ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COLLATE=utf8mb4_unicode_ci
    """)

    # 6. edit_requests
    query_run("""
        CREATE TABLE IF NOT EXISTS edit_requests (
          id INT AUTO_INCREMENT PRIMARY KEY,
          student_id INT NOT NULL,
          coordinator_name VARCHAR(255) NOT NULL DEFAULT 'Coordinator',
          event_name VARCHAR(255) DEFAULT '',
          old_data TEXT NOT NULL,
          new_data TEXT NOT NULL,
          status VARCHAR(50) NOT NULL DEFAULT 'Pending',
          admin_name VARCHAR(255) DEFAULT 'Admin',
          created_at DATETIME DEFAULT CURRENT_TIMESTAMP,
          approved_at DATETIME DEFAULT NULL,
          INDEX idx_er_student_id (student_id),
          INDEX idx_er_status (status),
          CONSTRAINT fk_er_student FOREIGN KEY (student_id) REFERENCES students(id) ON DELETE CASCADE
        ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COLLATE=utf8mb4_unicode_ci
    """)

    # 7. event_change_requests
    query_run("""
        CREATE TABLE IF NOT EXISTS event_change_requests (
          id INT AUTO_INCREMENT PRIMARY KEY,
          student_id INT NOT NULL,
          coordinator_name VARCHAR(255) NOT NULL DEFAULT 'Coordinator',
          event_name VARCHAR(255) DEFAULT '',
          old_events TEXT NOT NULL,
          new_events TEXT NOT NULL,
          reason TEXT DEFAULT NULL,
          status VARCHAR(50) NOT NULL DEFAULT 'Pending',
          admin_name VARCHAR(255) DEFAULT 'Admin',
          created_at DATETIME DEFAULT CURRENT_TIMESTAMP,
          approved_at DATETIME DEFAULT NULL,
          INDEX idx_ecr_student_id (student_id),
          INDEX idx_ecr_status (status),
          CONSTRAINT fk_ecr_student FOREIGN KEY (student_id) REFERENCES students(id) ON DELETE CASCADE
        ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COLLATE=utf8mb4_unicode_ci
    """)

    # 8. event_teams
    query_run("""
        CREATE TABLE IF NOT EXISTS event_teams (
          id INT AUTO_INCREMENT PRIMARY KEY,
          team_code VARCHAR(100) UNIQUE NOT NULL,
          team_name VARCHAR(255) NOT NULL,
          event_id VARCHAR(50) NOT NULL,
          event_name VARCHAR(255) NOT NULL,
          status VARCHAR(50) DEFAULT 'Formed',
          created_at DATETIME DEFAULT CURRENT_TIMESTAMP,
          INDEX idx_et_event_id (event_id),
          CONSTRAINT fk_et_event FOREIGN KEY (event_id) REFERENCES events(id) ON DELETE CASCADE
        ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COLLATE=utf8mb4_unicode_ci
    """)

    # 9. event_team_members
    query_run("""
        CREATE TABLE IF NOT EXISTS event_team_members (
          id INT AUTO_INCREMENT PRIMARY KEY,
          team_id INT NOT NULL,
          student_id INT NOT NULL,
          event_id VARCHAR(50) NOT NULL,
          joined_at DATETIME DEFAULT CURRENT_TIMESTAMP,
          UNIQUE KEY uq_event_student (event_id, student_id),
          INDEX idx_etm_team_id (team_id),
          CONSTRAINT fk_etm_team FOREIGN KEY (team_id) REFERENCES event_teams(id) ON DELETE CASCADE,
          CONSTRAINT fk_etm_student FOREIGN KEY (student_id) REFERENCES students(id) ON DELETE CASCADE,
          CONSTRAINT fk_etm_event FOREIGN KEY (event_id) REFERENCES events(id) ON DELETE CASCADE
        ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COLLATE=utf8mb4_unicode_ci
    """)

    # 10. counters
    query_run("""
        CREATE TABLE IF NOT EXISTS counters (
          id INT AUTO_INCREMENT PRIMARY KEY,
          counter_name VARCHAR(100) UNIQUE NOT NULL,
          faculty_name VARCHAR(255) DEFAULT NULL,
          status VARCHAR(50) DEFAULT 'ACTIVE',
          updated_at DATETIME DEFAULT CURRENT_TIMESTAMP ON UPDATE CURRENT_TIMESTAMP,
          INDEX idx_counters_name (counter_name)
        ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COLLATE=utf8mb4_unicode_ci
    """)

    # 11. faculty_members
    query_run("""
        CREATE TABLE IF NOT EXISTS faculty_members (
          id INT AUTO_INCREMENT PRIMARY KEY,
          name VARCHAR(255) UNIQUE NOT NULL,
          department VARCHAR(100) DEFAULT 'Academics',
          created_at DATETIME DEFAULT CURRENT_TIMESTAMP
        ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COLLATE=utf8mb4_unicode_ci
    """)

    # Seed events
    for evt in EVENTS:
        query_run(
            "REPLACE INTO events (id, name, day, category, is_standalone, min_participants, max_participants) VALUES (?, ?, ?, ?, ?, ?, ?)",
            (evt['id'], evt['name'], evt['day'], evt['category'], 1 if evt.get('is_standalone') else 0, evt.get('min_participants', 1), evt.get('max_participants', 1))
        )

    # Pre-seed 10 counters
    for i in range(1, 11):
        c_name = f"Counter {i}"
        query_run("INSERT IGNORE INTO counters (counter_name, faculty_name, status) VALUES (?, NULL, 'ACTIVE')", (c_name,))

    # Pre-seed faculty
    default_faculty = [
        'Dr. Kumar', 'Mrs. Priya', 'Mr. Arun', 'Dr. Ramesh',
        'Prof. Anita', 'Dr. Suresh', 'Mrs. Deepa', 'Mr. Karthik',
        'Dr. Meenakshi', 'Prof. Rajesh', 'Dr. Saravanan', 'Mrs. Lakshmi',
        'Mr. Balaji', 'Prof. Anitha', 'Dr. Venkatesh', 'Mrs. Gayathri'
    ]
    for fac in default_faculty:
        query_run("INSERT IGNORE INTO faculty_members (name) VALUES (?)", (fac,))


# ==========================================
# STATIC FRONTEND ROUTES
# ==========================================
@app.route('/')
def serve_index():
    return send_from_directory('public', 'index.html')

@app.route('/<path:path>')
def serve_static_files(path):
    if os.path.exists(os.path.join('public', path)):
        return send_from_directory('public', path)
    return send_from_directory('public', 'index.html')


# ==========================================
# 0. REGISTRATION & DASHBOARD STATS
# ==========================================
def get_stats_data():
    total_row = query_get("SELECT COUNT(*) as total FROM students")
    online_row = query_get("SELECT COUNT(*) as count FROM students WHERE UPPER(registration_type) = 'ONLINE' OR registration_type IS NULL")
    spot_row = query_get("SELECT COUNT(*) as count FROM students WHERE UPPER(registration_type) = 'SPOT'")
    day1_row = query_get("SELECT COUNT(*) as count FROM students WHERE day_selection = 'day1'")
    day2_row = query_get("SELECT COUNT(*) as count FROM students WHERE day_selection = 'day2'")
    both_row = query_get("SELECT COUNT(*) as count FROM students WHERE day_selection = 'both'")
    revenue_row = query_get("SELECT SUM(total_fee) as total FROM students")

    ff_total_players_row = query_get("SELECT COUNT(*) as count FROM freefire_players")
    ff_online_solo_row = query_get("SELECT COUNT(*) as count FROM freefire_players WHERE UPPER(registration_type) = 'ONLINE' AND team_id IS NULL")
    ff_spot_players_row = query_get("SELECT COUNT(*) as count FROM freefire_players WHERE UPPER(registration_type) = 'SPOT'")
    ff_teams_row = query_get("SELECT COUNT(*) as count FROM freefire_teams")

    pending_profile_reqs_row = query_get("SELECT COUNT(*) as count FROM edit_requests WHERE UPPER(status) = 'PENDING'")
    pending_event_reqs_row = query_get("SELECT COUNT(*) as count FROM event_change_requests WHERE UPPER(status) = 'PENDING'")

    return {
        'totalStudents': total_row['total'] if total_row else 0,
        'onlineCount': online_row['count'] if online_row else 0,
        'spotCount': spot_row['count'] if spot_row else 0,
        'day1Count': day1_row['count'] if day1_row else 0,
        'day2Count': day2_row['count'] if day2_row else 0,
        'bothDaysCount': both_row['count'] if both_row else 0,
        'totalRevenue': revenue_row['total'] if revenue_row and revenue_row['total'] is not None else 0,
        'freefireTotalPlayers': ff_total_players_row['count'] if ff_total_players_row else 0,
        'freefireOnlineSolo': ff_online_solo_row['count'] if ff_online_solo_row else 0,
        'freefireSpotPlayers': ff_spot_players_row['count'] if ff_spot_players_row else 0,
        'freefireTeamsCount': ff_teams_row['count'] if ff_teams_row else 0,
        'pendingProfileRequestsCount': pending_profile_reqs_row['count'] if pending_profile_reqs_row else 0,
        'pendingEventRequestsCount': pending_event_reqs_row['count'] if pending_event_reqs_row else 0
    }

@app.route('/api/stats', methods=['GET'])
def get_stats():
    try:
        stats = get_stats_data()
        return jsonify({'success': True, 'stats': stats})
    except Exception as e:
        print(f"Error fetching stats: {e}")
        return jsonify({'success': False, 'message': 'Failed to fetch statistics'}), 500

@app.route('/api/admin/stats-summary', methods=['GET'])
def get_admin_stats_summary():
    try:
        stats = get_stats_data()
        return jsonify({'success': True, 'stats': stats})
    except Exception as e:
        print(f"Error fetching admin stats summary: {e}")
        return jsonify({'success': False, 'message': 'Failed to fetch admin stats summary'}), 500


# ==========================================
# 1. EVENT CATALOG
# ==========================================
@app.route('/api/events', methods=['GET'])
def get_events():
    try:
        counts = query_all("SELECT event_id, COUNT(*) as count FROM student_events GROUP BY event_id")
        count_map = {row['event_id']: row['count'] for row in counts}

        events_with_counts = []
        for evt in EVENTS:
            item = dict(evt)
            item['registered_count'] = count_map.get(evt['id'], 0)
            events_with_counts.append(item)

        return jsonify({'success': True, 'events': events_with_counts})
    except Exception as e:
        print(f"Error fetching events: {e}")
        return jsonify({'success': False, 'message': 'Failed to fetch events'}), 500


# ==========================================
# 2. SPOT REGISTRATION
# ==========================================
@app.route('/api/register', methods=['POST'])
def register_student():
    try:
        data = request.get_json() or {}
        name = data.get('name', '').strip()
        phone = data.get('phone', '').strip()
        email = data.get('email', '').strip()
        college = data.get('college', '').strip()
        day_selection = data.get('daySelection', '')
        selected_event_ids = data.get('selectedEventIds') or data.get('events') or data.get('selectedEvents') or data.get('eventIds') or []
        hackathon_theme = data.get('hackathonTheme')
        counter_name = data.get('counterName')
        faculty_name = data.get('facultyName')

        if not name:
            return jsonify({'success': False, 'message': 'Student Name is required.'}), 400
        if not phone or not re.match(r'^\d{10}$', phone):
            return jsonify({'success': False, 'message': 'Valid 10-digit Phone Number is required.'}), 400
        if not email:
            return jsonify({'success': False, 'message': 'Email Address is required.'}), 400
        if not college:
            return jsonify({'success': False, 'message': 'College Name is required.'}), 400
        if day_selection not in ['day1', 'day2', 'both']:
            return jsonify({'success': False, 'message': 'Invalid Day Selection.'}), 400
        if not isinstance(selected_event_ids, list) or len(selected_event_ids) == 0:
            return jsonify({'success': False, 'message': 'At least one event must be selected.'}), 400

        selected_events = [e for e in EVENTS if e['id'] in selected_event_ids]
        if len(selected_events) != len(selected_event_ids):
            return jsonify({'success': False, 'message': 'One or more invalid events selected.'}), 400

        is_freefire_selected = 'd1_freefire' in selected_event_ids
        if is_freefire_selected and len(selected_event_ids) > 1:
            return jsonify({
                'success': False,
                'message': 'Free Fire Tournament is a standalone event. You cannot combine it with other events in the standard pass.'
            }), 400
        if day_selection == 'both' and is_freefire_selected:
            return jsonify({
                'success': False,
                'message': 'Free Fire Tournament is not available under Both Days pass.'
            }), 400

        day1_selected = [e for e in selected_events if e['day'] == 'day1']
        day2_selected = [e for e in selected_events if e['day'] == 'day2']

        if day_selection == 'day1':
            if len(day2_selected) > 0:
                return jsonify({'success': False, 'message': 'Day 1 registration cannot include Day 2 events.'}), 400
            if len(day1_selected) > LIMITS['DAY1_MAX_EVENTS']:
                return jsonify({'success': False, 'message': f"Maximum {LIMITS['DAY1_MAX_EVENTS']} events allowed for Day 1."}), 400
        elif day_selection == 'day2':
            if len(day1_selected) > 0:
                return jsonify({'success': False, 'message': 'Day 2 registration cannot include Day 1 events.'}), 400
            if len(day2_selected) > LIMITS['DAY2_MAX_EVENTS']:
                return jsonify({'success': False, 'message': f"Maximum {LIMITS['DAY2_MAX_EVENTS']} events allowed for Day 2."}), 400
        elif day_selection == 'both':
            if len(day1_selected) > LIMITS['DAY1_MAX_EVENTS']:
                return jsonify({'success': False, 'message': f"Maximum {LIMITS['DAY1_MAX_EVENTS']} events allowed for Day 1."}), 400
            if len(day2_selected) > LIMITS['DAY2_MAX_EVENTS']:
                return jsonify({'success': False, 'message': f"Maximum {LIMITS['DAY2_MAX_EVENTS']} events allowed for Day 2."}), 400
            if len(selected_events) > LIMITS['BOTH_DAYS_MAX_EVENTS']:
                return jsonify({'success': False, 'message': f"Maximum {LIMITS['BOTH_DAYS_MAX_EVENTS']} events total allowed for Both Days."}), 400

        if is_freefire_selected:
            total_fee = PRICING['FREE_FIRE_STANDALONE']
        elif day_selection == 'both':
            total_fee = PRICING['BOTH_DAYS']
        elif day_selection == 'day1':
            total_fee = PRICING['DAY1_ONLY']
        else:
            total_fee = PRICING['DAY2_ONLY']

        random_suffix = random.randint(1000, 9999)
        reg_code = f"SPOT-2026-{random_suffix}"

        final_counter = str(counter_name).strip() if counter_name and str(counter_name).strip() else None
        final_faculty = str(faculty_name).strip() if faculty_name and str(faculty_name).strip() else None

        if final_counter and not final_faculty:
            c_row = query_get("SELECT faculty_name FROM counters WHERE counter_name = ?", (final_counter,))
            if c_row and c_row.get('faculty_name'):
                final_faculty = c_row['faculty_name']

        conn = get_connection()
        try:
            conn.begin()
            with conn.cursor() as cur:
                cur.execute(
                    "INSERT INTO students (reg_code, name, phone, email, college, day_selection, total_fee, registration_type, counter_name, faculty_name) VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s)",
                    (reg_code, name, phone, email, college, day_selection, total_fee, 'SPOT', final_counter, final_faculty)
                )
                student_id = cur.lastrowid

                for evt in selected_events:
                    theme_val = hackathon_theme or 'AI & Smart Automation' if evt['id'] == 'd1_hackathon' else None
                    cur.execute(
                        "INSERT INTO student_events (student_id, event_id, event_day, verified_status, hackathon_theme) VALUES (%s, %s, %s, %s, %s)",
                        (student_id, evt['id'], evt['day'], 0, theme_val)
                    )

                if is_freefire_selected:
                    cur.execute(
                        "INSERT INTO freefire_players (team_id, name, phone, email, college, registration_type, is_captain, counter_name, faculty_name) VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s)",
                        (None, name, phone, email, college, 'SPOT', 0, final_counter, final_faculty)
                    )

            conn.commit()
            return jsonify({
                'success': True,
                'message': 'Spot Registration completed successfully!',
                'studentId': student_id,
                'regCode': reg_code,
                'registration': {
                    'id': student_id,
                    'student_id': student_id,
                    'reg_code': reg_code,
                    'regCode': reg_code,
                    'name': name,
                    'phone': phone,
                    'email': email,
                    'college': college,
                    'day_selection': day_selection,
                    'total_fee': total_fee,
                    'selected_events': selected_events,
                    'hackathon_theme': hackathon_theme,
                    'counter_name': final_counter,
                    'faculty_name': final_faculty
                }
            })
        except Exception as tx_err:
            conn.rollback()
            raise tx_err
        finally:
            conn.close()

    except Exception as e:
        print(f"Registration error: {e}")
        return jsonify({'success': False, 'message': 'Server error processing spot registration.'}), 500


# ==========================================
# 3. GET REGISTRATIONS (Search & Filter)
# ==========================================
@app.route('/api/registrations', methods=['GET'])
def get_registrations():
    try:
        q = (request.args.get('q') or request.args.get('search') or '').strip()
        event_id = request.args.get('eventId')
        day = request.args.get('day')

        sql = """
          SELECT 
            s.id, s.reg_code, s.name, s.phone, s.college, s.email, s.day_selection, s.total_fee, s.food_given, s.tag_given, s.food_d2_given, s.tag_d2_given, s.registered_at, s.registration_type,
            se.event_id, se.event_day, se.verified_status, se.hackathon_theme, se.event_status, se.completed_at,
            e.name as event_name, e.category as event_category,
            et.id as team_id, et.team_code, et.team_name
          FROM students s
          LEFT JOIN student_events se ON s.id = se.student_id
          LEFT JOIN events e ON se.event_id = e.id
          LEFT JOIN event_team_members etm ON (etm.student_id = s.id AND etm.event_id = se.event_id)
          LEFT JOIN event_teams et ON etm.team_id = et.id
          WHERE 1=1
        """
        params = []

        if event_id:
            sql += " AND se.event_id = ?"
            params.append(event_id)

        if day:
            if day == 'day1':
                sql += " AND (s.day_selection = 'day1' OR s.day_selection = 'both')"
            elif day == 'day2':
                sql += " AND (s.day_selection = 'day2' OR s.day_selection = 'both')"
            elif day == 'both':
                sql += " AND s.day_selection = 'both'"

        if q:
            term = f"%{q}%"
            sql += " AND (s.name LIKE ? OR s.phone LIKE ? OR s.college LIKE ? OR s.email LIKE ? OR s.reg_code LIKE ?)"
            params.extend([term, term, term, term, term])

        sql += " ORDER BY s.registered_at DESC, s.id DESC"

        rows = query_all(sql, params)

        student_map = {}
        for r in rows:
            sid = r['id']
            if sid not in student_map:
                student_map[sid] = {
                    'id': r['id'],
                    'reg_code': r['reg_code'],
                    'name': r['name'],
                    'phone': r['phone'],
                    'college': r['college'],
                    'email': r.get('email') or '',
                    'day_selection': r['day_selection'],
                    'total_fee': r['total_fee'],
                    'registration_type': (r.get('registration_type') or 'ONLINE').upper(),
                    'food_given': bool(r['food_given']),
                    'tag_given': bool(r['tag_given']),
                    'food_d2_given': bool(r.get('food_d2_given')),
                    'tag_d2_given': bool(r.get('tag_d2_given')),
                    'registered_at': r['registered_at'],
                    'team_id': r.get('team_id'),
                    'team_code': r.get('team_code'),
                    'team_name': r.get('team_name'),
                    'events': []
                }
            if r.get('event_id'):
                student_map[sid]['events'].append({
                    'event_id': r['event_id'],
                    'event_name': r['event_name'],
                    'event_day': r['event_day'],
                    'event_category': r['event_category'],
                    'verified_status': bool(r['verified_status']),
                    'hackathon_theme': r.get('hackathon_theme'),
                    'event_status': r.get('event_status') or 'NOT_STARTED',
                    'completed_at': r.get('completed_at'),
                    'team_id': r.get('team_id'),
                    'team_code': r.get('team_code'),
                    'team_name': r.get('team_name')
                })

        students = list(student_map.values())
        return jsonify({'success': True, 'count': len(students), 'students': students, 'registrations': students})
    except Exception as e:
        print(f"Error fetching registrations: {e}")
        return jsonify({'success': False, 'message': 'Failed to fetch registrations.'}), 500


# ==========================================
# 4. COORDINATOR EVENT TEAMS & ACTIONS
# ==========================================
@app.route('/api/coordinator/event-teams', methods=['GET'])
def get_coordinator_event_teams():
    try:
        event_id = request.args.get('eventId')
        if not event_id:
            return jsonify({'success': False, 'message': 'eventId query parameter is required'}), 400

        teams = query_all("SELECT id, team_code, team_name, event_id, event_name, status, created_at FROM event_teams WHERE event_id = ? ORDER BY id ASC", (event_id,))
        team_ids = [t['id'] for t in teams]

        members_by_team = {}
        if team_ids:
            placeholders = ','.join(['?'] * len(team_ids))
            members = query_all(f"""
                SELECT etm.team_id, s.id as student_id, s.name, s.phone, s.college, s.email, s.reg_code, etm.joined_at
                FROM event_team_members etm
                JOIN students s ON etm.student_id = s.id
                WHERE etm.team_id IN ({placeholders})
                ORDER BY etm.id ASC
            """, team_ids)
            for m in members:
                members_by_team.setdefault(m['team_id'], []).append(m)

        teams_with_members = []
        for t in teams:
            m_list = members_by_team.get(t['id'], [])
            t_obj = dict(t)
            t_obj['members_count'] = len(m_list)
            t_obj['members'] = m_list
            teams_with_members.append(t_obj)

        return jsonify({'success': True, 'count': len(teams_with_members), 'teams': teams_with_members})
    except Exception as e:
        print(f"Error fetching event teams: {e}")
        return jsonify({'success': False, 'message': 'Failed to fetch event teams'}), 500

@app.route('/api/coordinator/create-team', methods=['POST'])
def create_coordinator_team():
    try:
        data = request.get_json() or {}
        event_id = data.get('eventId')
        team_name = (data.get('teamName') or '').strip()
        student_ids = data.get('studentIds', [])

        if not event_id:
            return jsonify({'success': False, 'message': 'eventId is required'}), 400
        if not isinstance(student_ids, list) or len(student_ids) == 0:
            return jsonify({'success': False, 'message': 'Please select at least one participant.'}), 400

        event = next((e for e in EVENTS if e['id'] == event_id), None)
        if not event:
            return jsonify({'success': False, 'message': 'Event not found'}), 404

        min_p = event.get('min_participants', 1)
        max_p = event.get('max_participants', 1)

        if max_p <= 1:
            return jsonify({'success': False, 'message': f"Team formation is not applicable for individual/solo events ({event['name']})."}), 400
        if len(student_ids) < min_p:
            return jsonify({'success': False, 'message': f"Minimum {min_p} participant{'s' if min_p > 1 else ''} required for {event['name']}. You selected {len(student_ids)}."}), 400
        if len(student_ids) > max_p:
            return jsonify({'success': False, 'message': f"Maximum {max_p} participant{'s' if max_p > 1 else ''} allowed for {event['name']}. You selected {len(student_ids)}."}), 400

        placeholders = ','.join(['?'] * len(student_ids))
        existing = query_all(f"""
            SELECT etm.student_id, s.name, et.team_code, et.team_name
            FROM event_team_members etm
            JOIN event_teams et ON etm.team_id = et.id
            JOIN students s ON etm.student_id = s.id
            WHERE etm.event_id = ? AND etm.student_id IN ({placeholders})
        """, [event_id] + student_ids)

        if existing:
            names = ', '.join([f"{m['name']} ({m['team_code']})" for m in existing])
            return jsonify({'success': False, 'message': f"The following participant(s) are already assigned to a team for {event['name']}: {names}"}), 400

        reg_check = query_all(f"SELECT student_id FROM student_events WHERE event_id = ? AND student_id IN ({placeholders})", [event_id] + student_ids)
        if len(reg_check) != len(student_ids):
            return jsonify({'success': False, 'message': 'One or more selected participants are not registered for this event.'}), 400

        count_row = query_get("SELECT COUNT(*) as count FROM event_teams WHERE event_id = ?", (event_id,))
        next_num = (count_row['count'] if count_row else 0) + 1
        team_code = f"TEAM-{str(next_num).zfill(3)}"
        final_team_name = team_name if team_name else f"Team {str(next_num).zfill(3)}"

        conn = get_connection()
        try:
            conn.begin()
            with conn.cursor() as cur:
                cur.execute(
                    "INSERT INTO event_teams (team_code, team_name, event_id, event_name, status) VALUES (%s, %s, %s, %s, %s)",
                    (team_code, final_team_name, event_id, event['name'], 'Formed')
                )
                team_id = cur.lastrowid
                for sid in student_ids:
                    cur.execute(
                        "INSERT INTO event_team_members (team_id, student_id, event_id) VALUES (%s, %s, %s)",
                        (team_id, sid, event_id)
                    )
            conn.commit()

            created_members = query_all(
                "SELECT s.id as student_id, s.name, s.phone, s.college, s.email, s.reg_code FROM event_team_members etm JOIN students s ON etm.student_id = s.id WHERE etm.team_id = ?",
                (team_id,)
            )

            return jsonify({
                'success': True,
                'message': f"Team {team_code} created successfully with {len(student_ids)} members!",
                'team': {
                    'id': team_id,
                    'team_code': team_code,
                    'team_name': final_team_name,
                    'event_id': event_id,
                    'event_name': event['name'],
                    'members_count': len(student_ids),
                    'members': created_members
                }
            })
        except Exception as tx_err:
            conn.rollback()
            raise tx_err
        finally:
            conn.close()

    except Exception as e:
        print(f"Error creating event team: {e}")
        return jsonify({'success': False, 'message': 'Server error creating event team'}), 500

@app.route('/api/coordinator/teams/<int:team_id>', methods=['DELETE'])
def disband_coordinator_team(team_id):
    try:
        team = query_get("SELECT * FROM event_teams WHERE id = ?", (team_id,))
        if not team:
            return jsonify({'success': False, 'message': 'Team not found'}), 404

        conn = get_connection()
        try:
            conn.begin()
            with conn.cursor() as cur:
                cur.execute("DELETE FROM event_team_members WHERE team_id = %s", (team_id,))
                cur.execute("DELETE FROM event_teams WHERE id = %s", (team_id,))
            conn.commit()
            return jsonify({'success': True, 'message': f"Team {team['team_code']} disbanded successfully. Members are now available for selection."})
        except Exception as tx_err:
            conn.rollback()
            raise tx_err
        finally:
            conn.close()

    except Exception as e:
        print(f"Error disbanding team: {e}")
        return jsonify({'success': False, 'message': 'Failed to disband team'}), 500

@app.route('/api/registrations/<int:student_id>/verify', methods=['POST'])
def verify_event_attendance(student_id):
    try:
        data = request.get_json() or {}
        event_id = data.get('eventId')
        verified_status = bool(data.get('verifiedStatus'))

        if not student_id or not event_id:
            return jsonify({'success': False, 'message': 'Student ID and Event ID are required.'}), 400

        query_run("UPDATE student_events SET verified_status = ? WHERE student_id = ? AND event_id = ?",
                  (1 if verified_status else 0, student_id, event_id))

        return jsonify({'success': True, 'message': 'Event attendance verified successfully.'})
    except Exception as e:
        print(f"Verification error: {e}")
        return jsonify({'success': False, 'message': 'Server error updating verification.'}), 500

@app.route('/api/coordinator/edit-request', methods=['POST'])
def submit_edit_request():
    try:
        data = request.get_json() or {}
        student_id = data.get('studentId')
        coordinator_name = data.get('coordinatorName')
        event_name = data.get('eventName')
        new_data_dict = data.get('newData') or {}

        target_name = (data.get('newName') or new_data_dict.get('name') or '').strip()
        target_phone = (data.get('newPhone') or new_data_dict.get('phone') or '').strip()
        target_college = (data.get('newCollege') or new_data_dict.get('college') or '').strip()
        target_email = (data.get('newEmail') or new_data_dict.get('email') or '').strip()

        if not student_id:
            return jsonify({'success': False, 'message': 'studentId is required.'}), 400
        if not target_name or not target_phone or not target_college:
            return jsonify({'success': False, 'message': 'Name, Phone, and College Name are required.'}), 400

        current_student = query_get("SELECT name, phone, email, college FROM students WHERE id = ?", (student_id,))
        if not current_student:
            return jsonify({'success': False, 'message': 'Student not found.'}), 404

        final_event_name = (event_name or '').strip()
        if not final_event_name:
            evts = query_all("SELECT e.name FROM student_events se JOIN events e ON se.event_id = e.id WHERE se.student_id = ?", (student_id,))
            final_event_name = ', '.join([e['name'] for e in evts]) if evts else 'Event Desk'

        old_data = {
            'name': current_student['name'],
            'phone': current_student['phone'],
            'email': current_student.get('email') or '',
            'college': current_student['college']
        }
        target_data = {
            'name': target_name,
            'phone': target_phone,
            'college': target_college,
            'email': target_email
        }

        query_run(
            "INSERT INTO edit_requests (student_id, coordinator_name, event_name, old_data, new_data, status) VALUES (?, ?, ?, ?, ?, 'Pending')",
            (student_id, (coordinator_name or final_event_name).strip(), final_event_name, json.dumps(old_data), json.dumps(target_data))
        )

        return jsonify({'success': True, 'message': 'Edit request submitted to Admin Dashboard for approval.'})
    except Exception as e:
        print(f"Error submitting edit request: {e}")
        return jsonify({'success': False, 'message': 'Server error submitting edit request.'}), 500

@app.route('/api/coordinator/action', methods=['POST'])
def coordinator_action():
    try:
        data = request.get_json() or {}
        student_id = data.get('studentId')
        action_type = (data.get('actionType') or data.get('action') or '').upper()
        event_id = data.get('eventId')

        if not student_id:
            return jsonify({'success': False, 'message': 'studentId is required.'}), 400

        student = query_get("SELECT * FROM students WHERE id = ?", (student_id,))
        if not student:
            return jsonify({'success': False, 'message': 'Student record not found.'}), 404

        if action_type in ['FOOD', 'ISSUE_FOOD', 'ISSUE_FOOD_D1', 'FOOD_D1', 'FOOD_TOKEN']:
            if student['food_given'] == 1:
                return jsonify({'success': False, 'message': 'Day 1 Food token has already been issued.'}), 400
            query_run("UPDATE students SET food_given = 1 WHERE id = ?", (student_id,))
            return jsonify({'success': True, 'message': 'Day 1 Food Token marked as given successfully.'})

        if action_type in ['ISSUE_FOOD_D2', 'FOOD_D2']:
            if student.get('food_d2_given') == 1:
                return jsonify({'success': False, 'message': 'Day 2 Food token has already been issued.'}), 400
            query_run("UPDATE students SET food_d2_given = 1 WHERE id = ?", (student_id,))
            return jsonify({'success': True, 'message': 'Day 2 Food Token marked as given successfully.'})

        if action_type in ['TAG', 'ISSUE_TAG', 'ISSUE_TAG_D1', 'TAG_D1', 'TAG_GIVEN']:
            if student['tag_given'] == 1:
                return jsonify({'success': False, 'message': 'Day 1 Tag has already been issued.'}), 400
            query_run("UPDATE students SET tag_given = 1 WHERE id = ?", (student_id,))
            return jsonify({'success': True, 'message': 'Day 1 Tag marked as given successfully.'})

        if action_type in ['ISSUE_TAG_D2', 'TAG_D2']:
            if student.get('tag_d2_given') == 1:
                return jsonify({'success': False, 'message': 'Day 2 Tag has already been issued.'}), 400
            query_run("UPDATE students SET tag_d2_given = 1 WHERE id = ?", (student_id,))
            return jsonify({'success': True, 'message': 'Day 2 Tag marked as given successfully.'})

        if action_type in ['EVENT_ENTRY', 'ENTER_EVENT']:
            if not event_id:
                return jsonify({'success': False, 'message': 'eventId is required.'}), 400
            query_run("UPDATE student_events SET event_status = 'COMPLETED', completed_at = CURRENT_TIMESTAMP WHERE student_id = ? AND event_id = ?", (student_id, event_id))
            return jsonify({'success': True, 'message': 'Event entry confirmed & marked COMPLETED.'})

        if action_type in ['CONFIRM', 'CONFIRM_LOCK']:
            return jsonify({'success': True, 'message': 'Coordinator actions locked & saved successfully.'})

        return jsonify({'success': False, 'message': 'Invalid actionType.'}), 400
    except Exception as e:
        print(f"Error processing coordinator action: {e}")
        return jsonify({'success': False, 'message': 'Failed to process coordinator action.'}), 500


# ==========================================
# 5. ADMIN EDIT & EVENT CHANGE REQUESTS
# ==========================================
@app.route('/api/admin/edit-requests', methods=['GET'])
def get_admin_edit_requests():
    try:
        rows = query_all("""
          SELECT 
            er.id, er.student_id, er.coordinator_name, er.event_name, er.old_data, er.new_data, er.status, er.created_at, er.approved_at, er.admin_name,
            s.reg_code, s.name as current_name,
            (SELECT GROUP_CONCAT(e.name SEPARATOR ', ') FROM student_events se JOIN events e ON se.event_id = e.id WHERE se.student_id = er.student_id) as student_events_list
          FROM edit_requests er
          JOIN students s ON er.student_id = s.id
          ORDER BY er.created_at DESC
        """)

        requests = []
        for r in rows:
            old_obj = r['old_data']
            new_obj = r['new_data']
            if isinstance(old_obj, str):
                try: old_obj = json.loads(old_obj)
                except: pass
            if isinstance(new_obj, str):
                try: new_obj = json.loads(new_obj)
                except: pass

            event_name = r['event_name'] if r.get('event_name') and r['event_name'].strip() else (r.get('student_events_list') or 'Event Desk')
            requests.append({
                'id': r['id'],
                'student_id': r['student_id'],
                'reg_code': r['reg_code'],
                'current_name': r['current_name'],
                'coordinator_name': r['coordinator_name'],
                'event_name': event_name,
                'old_data': old_obj,
                'new_data': new_obj,
                'status': r['status'],
                'created_at': r['created_at'],
                'approved_at': r['approved_at'],
                'admin_name': r.get('admin_name') or 'Admin'
            })

        return jsonify({'success': True, 'count': len(requests), 'requests': requests})
    except Exception as e:
        print(f"Error fetching edit requests: {e}")
        return jsonify({'success': False, 'message': 'Failed to fetch edit requests.'}), 500

@app.route('/api/admin/edit-request/approve', methods=['POST'])
def approve_edit_request():
    try:
        data = request.get_json() or {}
        request_id = data.get('requestId')
        admin_name = (data.get('adminName') or 'Admin').strip()

        if not request_id:
            return jsonify({'success': False, 'message': 'requestId is required.'}), 400

        req_row = query_get("SELECT * FROM edit_requests WHERE id = ?", (request_id,))
        if not req_row:
            return jsonify({'success': False, 'message': 'Request not found.'}), 404
        if req_row['status'].upper() != 'PENDING':
            return jsonify({'success': False, 'message': f"Request is already {req_row['status']}."}), 400

        new_data = req_row['new_data']
        if isinstance(new_data, str):
            new_data = json.loads(new_data)

        conn = get_connection()
        try:
            conn.begin()
            with conn.cursor() as cur:
                cur.execute(
                    "UPDATE students SET name = %s, phone = %s, college = %s, email = %s WHERE id = %s",
                    (new_data['name'], new_data['phone'], new_data['college'], new_data.get('email', ''), req_row['student_id'])
                )
                cur.execute(
                    "UPDATE edit_requests SET status = 'Approved', approved_at = CURRENT_TIMESTAMP, admin_name = %s WHERE id = %s",
                    (admin_name, request_id)
                )
            conn.commit()
            return jsonify({'success': True, 'message': 'Modification request approved & participant record updated!'})
        except Exception as tx_err:
            conn.rollback()
            raise tx_err
        finally:
            conn.close()

    except Exception as e:
        print(f"Error approving request: {e}")
        return jsonify({'success': False, 'message': 'Failed to approve modification request.'}), 500

@app.route('/api/admin/edit-request/reject', methods=['POST'])
def reject_edit_request():
    try:
        data = request.get_json() or {}
        request_id = data.get('requestId')
        admin_name = (data.get('adminName') or 'Admin').strip()

        if not request_id:
            return jsonify({'success': False, 'message': 'requestId is required.'}), 400

        query_run(
            "UPDATE edit_requests SET status = 'Rejected', approved_at = CURRENT_TIMESTAMP, admin_name = ? WHERE id = ? AND UPPER(status) = 'PENDING'",
            (admin_name, request_id)
        )
        return jsonify({'success': True, 'message': 'Modification request rejected. Original details kept.'})
    except Exception as e:
        print(f"Error rejecting request: {e}")
        return jsonify({'success': False, 'message': 'Failed to reject request.'}), 500

@app.route('/api/coordinator/event-change-request', methods=['POST'])
def submit_event_change_request():
    try:
        data = request.get_json() or {}
        student_id = data.get('studentId')
        coordinator_name = data.get('coordinatorName')
        event_name = data.get('eventName')
        new_event_ids = data.get('newEventIds', [])
        hackathon_theme = data.get('hackathonTheme')
        reason = data.get('reason', '')

        if not student_id:
            return jsonify({'success': False, 'message': 'studentId is required.'}), 400
        if not isinstance(new_event_ids, list) or len(new_event_ids) == 0:
            return jsonify({'success': False, 'message': 'At least one new event must be selected.'}), 400

        student = query_get("SELECT * FROM students WHERE id = ?", (student_id,))
        if not student:
            return jsonify({'success': False, 'message': 'Student record not found.'}), 404

        pending_req = query_get("SELECT id FROM event_change_requests WHERE student_id = ? AND UPPER(status) = 'PENDING'", (student_id,))
        if pending_req:
            return jsonify({'success': False, 'message': 'A pending event modification request already exists for this participant. Please wait for Admin review.'}), 400

        valid_events = [e for e in EVENTS if e['id'] in new_event_ids]
        if len(valid_events) != len(new_event_ids):
            return jsonify({'success': False, 'message': 'One or more selected events are invalid.'}), 400

        day1_events = [e for e in valid_events if e['day'] == 'day1']
        day2_events = [e for e in valid_events if e['day'] == 'day2']

        if student['day_selection'] == 'day1':
            if len(day2_events) > 0:
                return jsonify({'success': False, 'message': 'Participant has Day 1 Pass only. Cannot select Day 2 events.'}), 400
            if len(day1_events) > LIMITS['DAY1_MAX_EVENTS']:
                return jsonify({'success': False, 'message': f"Maximum {LIMITS['DAY1_MAX_EVENTS']} Day 1 events allowed."}), 400
        elif student['day_selection'] == 'day2':
            if len(day1_events) > 0:
                return jsonify({'success': False, 'message': 'Participant has Day 2 Pass only. Cannot select Day 1 events.'}), 400
            if len(day2_events) > LIMITS['DAY2_MAX_EVENTS']:
                return jsonify({'success': False, 'message': f"Maximum {LIMITS['DAY2_MAX_EVENTS']} Day 2 events allowed."}), 400
        elif student['day_selection'] == 'both':
            if len(day1_events) > LIMITS['DAY1_MAX_EVENTS']:
                return jsonify({'success': False, 'message': f"Maximum {LIMITS['DAY1_MAX_EVENTS']} Day 1 events allowed."}), 400
            if len(day2_events) > LIMITS['DAY2_MAX_EVENTS']:
                return jsonify({'success': False, 'message': f"Maximum {LIMITS['DAY2_MAX_EVENTS']} Day 2 events allowed."}), 400
            if len(valid_events) > LIMITS['BOTH_DAYS_MAX_EVENTS']:
                return jsonify({'success': False, 'message': f"Maximum {LIMITS['BOTH_DAYS_MAX_EVENTS']} events total allowed."}), 400

        current_events_rows = query_all("""
            SELECT se.event_id, se.event_day, se.hackathon_theme, e.name as event_name, e.category
            FROM student_events se
            JOIN events e ON se.event_id = e.id
            WHERE se.student_id = ?
        """, (student_id,))

        old_events_data = [{
            'id': r['event_id'],
            'name': r['event_name'],
            'day': r['event_day'],
            'category': r['category'],
            'hackathon_theme': r.get('hackathon_theme')
        } for r in current_events_rows]

        new_events_data = [{
            'id': evt['id'],
            'name': evt['name'],
            'day': evt['day'],
            'category': evt['category'],
            'hackathon_theme': (hackathon_theme or 'AI & Smart Automation') if evt['id'] == 'd1_hackathon' else None
        } for evt in valid_events]

        final_event_name = (event_name or '').strip()
        if not final_event_name:
            final_event_name = ', '.join([e['event_name'] for e in current_events_rows]) if current_events_rows else 'Event Issue Desk'

        query_run(
            "INSERT INTO event_change_requests (student_id, coordinator_name, event_name, old_events, new_events, reason, status) VALUES (?, ?, ?, ?, ?, ?, 'Pending')",
            (student_id, (coordinator_name or final_event_name).strip(), final_event_name, json.dumps(old_events_data), json.dumps(new_events_data), (reason or '').strip())
        )

        return jsonify({'success': True, 'message': 'Event modification request submitted to Admin Dashboard for approval!'})
    except Exception as e:
        print(f"Error submitting event change request: {e}")
        return jsonify({'success': False, 'message': 'Server error submitting event change request.'}), 500

@app.route('/api/admin/event-change-requests', methods=['GET'])
def get_admin_event_change_requests():
    try:
        rows = query_all("""
          SELECT 
            ecr.id, ecr.student_id, ecr.coordinator_name, ecr.event_name, ecr.old_events, ecr.new_events, ecr.reason, ecr.status, ecr.created_at, ecr.approved_at, ecr.admin_name,
            s.reg_code, s.name as student_name, s.phone, s.college, s.day_selection,
            (SELECT GROUP_CONCAT(e.name SEPARATOR ', ') FROM student_events se JOIN events e ON se.event_id = e.id WHERE se.student_id = ecr.student_id) as student_events_list
          FROM event_change_requests ecr
          JOIN students s ON ecr.student_id = s.id
          ORDER BY ecr.created_at DESC
        """)

        requests = []
        for r in rows:
            old_evts = []
            new_evts = []
            try: old_evts = json.loads(r['old_events'] or '[]')
            except: pass
            try: new_evts = json.loads(r['new_events'] or '[]')
            except: pass

            event_name = r['event_name'] if r.get('event_name') and r['event_name'].strip() else (r.get('student_events_list') or 'Event Issue Desk')
            requests.append({
                'id': r['id'],
                'student_id': r['student_id'],
                'reg_code': r['reg_code'],
                'student_name': r['student_name'],
                'phone': r['phone'],
                'college': r['college'],
                'day_selection': r['day_selection'],
                'coordinator_name': r['coordinator_name'],
                'event_name': event_name,
                'old_events': old_evts,
                'new_events': new_evts,
                'reason': r.get('reason'),
                'status': r['status'],
                'created_at': r['created_at'],
                'approved_at': r['approved_at'],
                'admin_name': r.get('admin_name')
            })

        return jsonify({'success': True, 'count': len(requests), 'requests': requests})
    except Exception as e:
        print(f"Error fetching event change requests: {e}")
        return jsonify({'success': False, 'message': 'Failed to fetch event change requests.'}), 500

@app.route('/api/admin/event-change-request/approve', methods=['POST'])
def approve_event_change_request():
    try:
        data = request.get_json() or {}
        request_id = data.get('requestId')
        admin_name = (data.get('adminName') or 'Admin').strip()

        if not request_id:
            return jsonify({'success': False, 'message': 'requestId is required.'}), 400

        req_row = query_get("SELECT * FROM event_change_requests WHERE id = ?", (request_id,))
        if not req_row:
            return jsonify({'success': False, 'message': 'Request not found.'}), 404
        if req_row['status'].upper() != 'PENDING':
            return jsonify({'success': False, 'message': f"Request is already {req_row['status']}."}), 400

        new_events = []
        try:
            new_events = json.loads(req_row['new_events'] or '[]')
        except:
            new_events = []

        if not new_events:
            return jsonify({'success': False, 'message': 'No new events found in request.'}), 400

        conn = get_connection()
        try:
            conn.begin()
            with conn.cursor() as cur:
                cur.execute("DELETE FROM student_events WHERE student_id = %s", (req_row['student_id'],))
                for evt in new_events:
                    cur.execute(
                        "INSERT INTO student_events (student_id, event_id, event_day, verified_status, hackathon_theme, event_status) VALUES (%s, %s, %s, %s, %s, %s)",
                        (req_row['student_id'], evt['id'], evt['day'], 0, evt.get('hackathon_theme'), 'NOT_STARTED')
                    )
                cur.execute(
                    "UPDATE event_change_requests SET status = 'Approved', approved_at = CURRENT_TIMESTAMP, admin_name = %s WHERE id = %s",
                    (admin_name, request_id)
                )
            conn.commit()
            return jsonify({'success': True, 'message': 'Event modification approved! Participant event registrations updated successfully.'})
        except Exception as tx_err:
            conn.rollback()
            raise tx_err
        finally:
            conn.close()

    except Exception as e:
        print(f"Error approving event change request: {e}")
        return jsonify({'success': False, 'message': 'Failed to approve event change request.'}), 500

@app.route('/api/admin/event-change-request/reject', methods=['POST'])
def reject_event_change_request():
    try:
        data = request.get_json() or {}
        request_id = data.get('requestId')
        admin_name = (data.get('adminName') or 'Admin').strip()

        if not request_id:
            return jsonify({'success': False, 'message': 'requestId is required.'}), 400

        query_run(
            "UPDATE event_change_requests SET status = 'Rejected', approved_at = CURRENT_TIMESTAMP, admin_name = ? WHERE id = ? AND UPPER(status) = 'PENDING'",
            (admin_name, request_id)
        )
        return jsonify({'success': True, 'message': 'Event modification request rejected. Original events kept unchanged.'})
    except Exception as e:
        print(f"Error rejecting event change request: {e}")
        return jsonify({'success': False, 'message': 'Failed to reject event change request.'}), 500

@app.route('/api/coordinator/pending-event-requests', methods=['GET'])
def get_pending_event_student_ids():
    try:
        rows = query_all("SELECT DISTINCT student_id FROM event_change_requests WHERE UPPER(status) = 'PENDING'")
        student_ids = [r['student_id'] for r in rows]
        return jsonify({'success': True, 'pendingStudentIds': student_ids})
    except Exception as e:
        print(f"Error fetching pending student IDs: {e}")
        return jsonify({'success': False, 'message': 'Failed to fetch pending student IDs.'}), 500


# ==========================================
# 6. ADMIN MASTER STUDENT DIRECTORY
# ==========================================
@app.route('/api/admin/all-students', methods=['GET'])
def get_all_students_admin():
    try:
        q = request.args.get('q', '').strip()
        source = request.args.get('source')
        pass_type = request.args.get('pass')
        event_id = request.args.get('eventId')

        sql = """
          SELECT 
            s.id, s.reg_code, s.name, s.phone, s.college, s.email, s.day_selection, s.total_fee, s.food_given, s.tag_given, s.food_d2_given, s.tag_d2_given, s.registered_at, s.registration_type,
            se.event_id, se.event_day, se.verified_status, se.hackathon_theme, se.event_status, se.completed_at,
            e.name as event_name, e.category as event_category
          FROM students s
          LEFT JOIN student_events se ON s.id = se.student_id
          LEFT JOIN events e ON se.event_id = e.id
          WHERE 1=1
        """
        params = []

        if source and source != 'all':
            sql += " AND UPPER(COALESCE(s.registration_type, 'ONLINE')) = ?"
            params.append(source.upper())

        if pass_type and pass_type != 'all':
            if pass_type == 'day1':
                sql += " AND (s.day_selection = 'day1' OR s.day_selection = 'both')"
            elif pass_type == 'day2':
                sql += " AND (s.day_selection = 'day2' OR s.day_selection = 'both')"
            elif pass_type == 'both':
                sql += " AND s.day_selection = 'both'"

        if event_id:
            sql += " AND se.event_id = ?"
            params.append(event_id)

        if q:
            term = f"%{q}%"
            sql += " AND (s.name LIKE ? OR s.phone LIKE ? OR s.college LIKE ? OR s.email LIKE ? OR s.reg_code LIKE ?)"
            params.extend([term, term, term, term, term])

        sql += " ORDER BY s.registered_at DESC, s.id DESC"

        rows = query_all(sql, params)

        student_map = {}
        for r in rows:
            sid = r['id']
            if sid not in student_map:
                student_map[sid] = {
                    'id': r['id'],
                    'reg_code': r['reg_code'],
                    'name': r['name'],
                    'phone': r['phone'],
                    'college': r['college'],
                    'email': r.get('email') or '',
                    'day_selection': r['day_selection'],
                    'total_fee': r['total_fee'],
                    'registration_type': (r.get('registration_type') or 'ONLINE').upper(),
                    'food_given': bool(r['food_given']),
                    'tag_given': bool(r['tag_given']),
                    'food_d2_given': bool(r.get('food_d2_given')),
                    'tag_d2_given': bool(r.get('tag_d2_given')),
                    'registered_at': r['registered_at'],
                    'events': []
                }
            if r.get('event_id'):
                student_map[sid]['events'].append({
                    'event_id': r['event_id'],
                    'event_name': r['event_name'],
                    'event_day': r['event_day'],
                    'event_category': r['event_category'],
                    'verified_status': bool(r['verified_status']),
                    'hackathon_theme': r.get('hackathon_theme'),
                    'event_status': r.get('event_status') or 'NOT_STARTED',
                    'completed_at': r.get('completed_at')
                })

        students = list(student_map.values())
        return jsonify({'success': True, 'count': len(students), 'students': students})
    except Exception as e:
        print(f"Error fetching all students for admin: {e}")
        return jsonify({'success': False, 'message': 'Failed to fetch student directory.'}), 500


# ==========================================
# 7. EXCEL / CSV EXPORT ENDPOINTS
# ==========================================
@app.route('/api/admin/export/excel', methods=['GET'])
def export_master_excel():
    try:
        student_rows = query_all("""
          SELECT 
            s.id, s.reg_code, s.name, s.phone, s.email, s.college, s.day_selection, s.total_fee, s.registration_type,
            s.food_given, s.tag_given, s.food_d2_given, s.tag_d2_given, s.registered_at,
            se.event_id, se.event_day, se.verified_status, se.hackathon_theme, se.event_status,
            e.name as event_name, e.category as event_category
          FROM students s
          LEFT JOIN student_events se ON s.id = se.student_id
          LEFT JOIN events e ON se.event_id = e.id
          ORDER BY s.id ASC
        """)

        student_map = {}
        for r in student_rows:
            sid = r['id']
            if sid not in student_map:
                student_map[sid] = {
                    'reg_code': r['reg_code'],
                    'name': r['name'],
                    'phone': r['phone'],
                    'email': r.get('email') or 'N/A',
                    'college': r['college'],
                    'day_selection': (r['day_selection'] or '').upper(),
                    'total_fee': r['total_fee'],
                    'registration_type': (r.get('registration_type') or 'ONLINE').upper(),
                    'food_given': 'YES' if r['food_given'] == 1 else 'NO',
                    'tag_given': 'YES' if r['tag_given'] == 1 else 'NO',
                    'registered_at': str(r['registered_at']),
                    'events': []
                }
            if r.get('event_name'):
                lbl = f"{r['event_name']} ({(r.get('event_day') or '').upper()})"
                if r.get('hackathon_theme'):
                    lbl += f" [Theme: {r['hackathon_theme']}]"
                student_map[sid]['events'].append(lbl)

        all_students = list(student_map.values())

        wb = openpyxl.Workbook()
        ws_master = wb.active
        ws_master.title = "All_Participants"

        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="4F46E5", end_color="4F46E5", fill_type="solid")

        def write_sheet(ws, headers, data):
            ws.append(headers)
            for cell in ws[1]:
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal="center", vertical="center")
            for row in data:
                ws.append(row)
            for col in ws.columns:
                max_len = max(len(str(cell.value or '')) for cell in col)
                col_letter = get_column_letter(col[0].column)
                ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

        # 1. Master
        master_headers = ['S.No', 'Registration Code', 'Participant Name', 'Phone Number', 'Email Address', 'College / Institute', 'Registration Type', 'Pass Type', 'Registration Fee (INR)', 'Food Token Issued', 'ID Tag Issued', 'Registered Events List', 'Registered Timestamp']
        master_rows = [[i + 1, s['reg_code'], s['name'], s['phone'], s['email'], s['college'], s['registration_type'], s['day_selection'], s['total_fee'], s['food_given'], s['tag_given'], ' | '.join(s['events']) or 'None', s['registered_at']] for i, s in enumerate(all_students)]
        write_sheet(ws_master, master_headers, master_rows)

        # 2. Online
        ws_online = wb.create_sheet("Online_Registrations")
        online_data = [[i + 1, s['reg_code'], s['name'], s['phone'], s['email'], s['college'], s['day_selection'], s['total_fee'], ' | '.join(s['events']) or 'None', s['registered_at']] for i, s in enumerate([s for s in all_students if s['registration_type'] == 'ONLINE'])]
        write_sheet(ws_online, ['S.No', 'Registration Code', 'Participant Name', 'Phone Number', 'Email Address', 'College / Institute', 'Pass Type', 'Fee (INR)', 'Registered Events', 'Registration Date'], online_data)

        # 3. Spot
        ws_spot = wb.create_sheet("Spot_Registrations")
        spot_data = [[i + 1, s['reg_code'], s['name'], s['phone'], s['email'], s['college'], s['day_selection'], s['total_fee'], s['food_given'], s['tag_given'], ' | '.join(s['events']) or 'None', s['registered_at']] for i, s in enumerate([s for s in all_students if s['registration_type'] == 'SPOT'])]
        write_sheet(ws_spot, ['S.No', 'Registration Code', 'Participant Name', 'Phone Number', 'Email Address', 'College / Institute', 'Pass Type', 'Fee (INR)', 'Food Token', 'ID Tag', 'Registered Events', 'Registration Date'], spot_data)

        # 4. Free Fire
        ws_ff = wb.create_sheet("Free_Fire_Tournament")
        ff_rows = query_all("""
            SELECT 
              t.team_code, t.team_name, t.registration_source, t.registered_at,
              p.name as player_name, p.phone as player_phone, p.email as player_email, p.college as player_college, p.is_captain
            FROM freefire_teams t
            JOIN freefire_players p ON t.id = p.team_id
            ORDER BY t.id ASC, p.is_captain DESC, p.id ASC
        """)
        ff_teams_map = {}
        for r in ff_rows:
            tc = r['team_code']
            if tc not in ff_teams_map:
                ff_teams_map[tc] = {
                    'team_code': r['team_code'],
                    'team_name': r['team_name'],
                    'registration_source': r['registration_source'],
                    'registered_at': str(r['registered_at']),
                    'captain_name': '-', 'captain_phone': '-', 'captain_email': '-', 'captain_college': '-',
                    'players': []
                }
            if r['is_captain'] == 1:
                ff_teams_map[tc]['captain_name'] = r['player_name']
                ff_teams_map[tc]['captain_phone'] = r['player_phone']
                ff_teams_map[tc]['captain_email'] = r.get('player_email') or 'N/A'
                ff_teams_map[tc]['captain_college'] = r['player_college']
            ff_teams_map[tc]['players'].append(f"{r['player_name']} ({r['player_phone']}) - {r['player_college']}")

        ff_data = [[i + 1, t['team_code'], t['team_name'], t['registration_source'], t['captain_name'], t['captain_phone'], t['captain_email'], t['captain_college'], ' | '.join(t['players']), t['registered_at']] for i, t in enumerate(ff_teams_map.values())]
        write_sheet(ws_ff, ['S.No', 'Squad Code', 'Squad Name', 'Registration Source', 'Captain Name', 'Captain Phone', 'Captain Email', 'Captain College', '4-Player Squad Details', 'Registered Timestamp'], ff_data)

        # 5. Events Summary
        ws_events = wb.create_sheet("Events_Summary")
        event_counts = query_all("""
          SELECT e.id as event_id, e.name as event_name, e.day as event_day, e.category,
                 COUNT(se.id) as registered_count,
                 SUM(CASE WHEN se.verified_status = 1 THEN 1 ELSE 0 END) as verified_count,
                 SUM(CASE WHEN se.event_status = 'COMPLETED' THEN 1 ELSE 0 END) as completed_count
          FROM events e
          LEFT JOIN student_events se ON e.id = se.event_id
          GROUP BY e.id
          ORDER BY e.day ASC, e.name ASC
        """)
        events_data = [[i + 1, ec['event_id'], ec['event_name'], (ec.get('event_day') or '').upper(), ec['category'], ec['registered_count'] or 0, ec['verified_count'] or 0, ec['completed_count'] or 0] for i, ec in enumerate(event_counts)]
        write_sheet(ws_events, ['S.No', 'Event ID', 'Event Name', 'Day', 'Category', 'Total Registered Students', 'Gate / Verified Count', 'Completed / In Event Count'], events_data)

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return send_file(buf, as_attachment=True, download_name="TechRaga26_Master_Registrations.xlsx", mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    except Exception as e:
        print(f"Error generating Excel export: {e}")
        return jsonify({'success': False, 'message': 'Failed to generate Excel sheet.'}), 500

@app.route('/api/admin/export/event/<event_id>', methods=['GET'])
def export_event_roster(event_id):
    try:
        evt = next((e for e in EVENTS if e['id'] == event_id), None)
        if not evt:
            return jsonify({'success': False, 'message': 'Event not found in catalog.'}), 404

        rows = query_all("""
          SELECT 
            s.id, s.reg_code, s.name, s.phone, s.email, s.college, s.day_selection, s.total_fee, s.registration_type,
            s.food_given, s.tag_given, s.registered_at,
            se.verified_status, se.hackathon_theme, se.event_status, se.completed_at
          FROM student_events se
          JOIN students s ON se.student_id = s.id
          WHERE se.event_id = ?
          ORDER BY s.name ASC
        """, (event_id,))

        wb = openpyxl.Workbook()
        ws = wb.active
        clean_sheet_name = re.sub(r'[\/\\?*:[\]]', '_', evt['name'])[:31]
        ws.title = clean_sheet_name

        headers = ['S.No', 'Registration Code', 'Participant Name', 'Phone Number', 'Email Address', 'College / Institute', 'Registration Type', 'Pass Type', 'Gate Attendance', 'Event Status ("In")', 'Food Token Issued', 'ID Tag Issued']
        if evt['id'] == 'd1_hackathon':
            headers.append('Hackathon Theme')
        headers.append('Registration Date')

        data = []
        for idx, r in enumerate(rows):
            row = [
                idx + 1, r['reg_code'], r['name'], r['phone'], r.get('email') or 'N/A', r['college'],
                (r.get('registration_type') or 'ONLINE').upper(), (r.get('day_selection') or '').upper(),
                'VERIFIED' if r['verified_status'] == 1 else 'PENDING',
                'COMPLETED (IN)' if r.get('event_status') == 'COMPLETED' else 'NOT STARTED',
                'YES' if r['food_given'] == 1 else 'NO',
                'YES' if r['tag_given'] == 1 else 'NO'
            ]
            if evt['id'] == 'd1_hackathon':
                row.append(r.get('hackathon_theme') or 'General')
            row.append(str(r['registered_at']))
            data.append(row)

        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="4F46E5", end_color="4F46E5", fill_type="solid")
        ws.append(headers)
        for cell in ws[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for row in data:
            ws.append(row)
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        clean_filename = f"{re.sub(r'[^a-zA-Z0-9_-]', '_', evt['name'])}_Roster.xlsx"
        return send_file(buf, as_attachment=True, download_name=clean_filename, mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    except Exception as e:
        print(f"Error exporting event roster: {e}")
        return jsonify({'success': False, 'message': 'Failed to export event roster.'}), 500

@app.route('/api/admin/export/csv', methods=['GET'])
def export_csv():
    try:
        req_type = request.args.get('type')
        sql = """
          SELECT 
            s.id, s.reg_code, s.name, s.phone, s.email, s.college, s.day_selection, s.total_fee, s.registration_type,
            s.food_given, s.tag_given, s.registered_at,
            e.name as event_name
          FROM students s
          LEFT JOIN student_events se ON s.id = se.student_id
          LEFT JOIN events e ON se.event_id = e.id
        """
        params = []
        if req_type == 'online':
            sql += " WHERE UPPER(COALESCE(s.registration_type, 'ONLINE')) = 'ONLINE'"
        elif req_type == 'spot':
            sql += " WHERE UPPER(COALESCE(s.registration_type, 'ONLINE')) = 'SPOT'"
        sql += " ORDER BY s.id ASC"

        rows = query_all(sql, params)

        student_map = {}
        for r in rows:
            sid = r['id']
            if sid not in student_map:
                student_map[sid] = {
                    'reg_code': r['reg_code'],
                    'name': r['name'],
                    'phone': r['phone'],
                    'email': r.get('email') or '',
                    'college': r['college'],
                    'registration_type': (r.get('registration_type') or 'ONLINE').upper(),
                    'day_selection': (r['day_selection'] or '').upper(),
                    'total_fee': r['total_fee'],
                    'food_given': 'Yes' if r['food_given'] == 1 else 'No',
                    'tag_given': 'Yes' if r['tag_given'] == 1 else 'No',
                    'registered_at': str(r['registered_at']),
                    'events': []
                }
            if r.get('event_name'):
                student_map[sid]['events'].append(r['event_name'])

        students = list(student_map.values())
        headers = ['S.No', 'Registration Code', 'Full Name', 'Phone', 'Email', 'College', 'Source', 'Pass', 'Fee Paid (INR)', 'Food Token', 'ID Tag', 'Events', 'Registered At']
        csv_lines = [','.join(headers)]

        for idx, s in enumerate(students):
            clean_name = str(s['name']).replace('"', '""')
            clean_email = str(s['email']).replace('"', '""')
            clean_college = str(s['college']).replace('"', '""')
            clean_events = ("; ".join(s['events'])).replace('"', '""')

            line = [
                str(idx + 1),
                f'"{s["reg_code"]}"',
                f'"{clean_name}"',
                f'"{s["phone"]}"',
                f'"{clean_email}"',
                f'"{clean_college}"',
                f'"{s["registration_type"]}"',
                f'"{s["day_selection"]}"',
                str(s['total_fee']),
                f'"{s["food_given"]}"',
                f'"{s["tag_given"]}"',
                f'"{clean_events}"',
                f'"{s["registered_at"]}"'
            ]
            csv_lines.append(','.join(line))

        output = '\r\n'.join(csv_lines)
        buf = io.BytesIO(output.encode('utf-8'))
        return send_file(buf, as_attachment=True, download_name=f"TechRaga26_Registrations_{req_type or 'all'}.csv", mimetype="text/csv")
    except Exception as e:
        print(f"Error generating CSV export: {e}")
        return jsonify({'success': False, 'message': 'Failed to generate CSV export.'}), 500


# ==========================================
# 8. ADMIN SYSTEM INFO
# ==========================================
@app.route('/api/admin/system-info', methods=['GET'])
def get_system_info():
    try:
        students_count = query_get("SELECT COUNT(*) as count FROM students")
        student_events_count = query_get("SELECT COUNT(*) as count FROM student_events")
        ff_players_count = query_get("SELECT COUNT(*) as count FROM freefire_players")
        ff_teams_count = query_get("SELECT COUNT(*) as count FROM freefire_teams")
        edit_reqs_count = query_get("SELECT COUNT(*) as count FROM edit_requests")
        event_change_reqs_count = query_get("SELECT COUNT(*) as count FROM event_change_requests")
        events_count = query_get("SELECT COUNT(*) as count FROM events")

        return jsonify({
            'success': True,
            'system': {
                'databaseEngine': 'MySQL (InnoDB)',
                'databaseFile': DB_NAME,
                'journalMode': 'InnoDB',
                'syncMode': 'ACID',
                'tables': {
                    'students': students_count['count'] if students_count else 0,
                    'student_events': student_events_count['count'] if student_events_count else 0,
                    'freefire_players': ff_players_count['count'] if ff_players_count else 0,
                    'freefire_teams': ff_teams_count['count'] if ff_teams_count else 0,
                    'edit_requests': edit_reqs_count['count'] if edit_reqs_count else 0,
                    'event_change_requests': event_change_reqs_count['count'] if event_change_reqs_count else 0,
                    'events': events_count['count'] if events_count else 0
                },
                'uptimeSeconds': int(os.times().elapsed),
                'serverTime': datetime.utcnow().isoformat()
            }
        })
    except Exception as e:
        print(f"Error fetching system info: {e}")
        return jsonify({'success': False, 'message': 'Failed to fetch system info.'}), 500


# ==========================================
# 9. EXCEL IMPORT ENGINE
# ==========================================
def match_event(raw_name):
    if not raw_name: return None
    clean = re.sub(r'[^a-z0-9]', '', str(raw_name).strip().lower())
    if not clean: return None

    for evt in EVENTS:
        evt_clean = re.sub(r'[^a-z0-9]', '', evt['name'].lower())
        id_clean = re.sub(r'[^a-z0-9]', '', evt['id'].lower())
        if clean == evt_clean or clean == id_clean:
            return evt

    if 'adzap' in clean or 'adz' in clean: return next(e for e in EVENTS if e['id'] == 'd1_adzap')
    if 'sharktank' in clean or 'businessplan' in clean or 'shark' in clean: return next(e for e in EVENTS if e['id'] == 'd1_sharktank')
    if 'quiz' in clean: return next(e for e in EVENTS if e['id'] == 'd1_quiz')
    if 'designer' in clean: return next(e for e in EVENTS if e['id'] == 'd1_designer')
    if 'fashion' in clean: return next(e for e in EVENTS if e['id'] == 'd1_fashion')
    if 'nailart' in clean or 'nail' in clean: return next(e for e in EVENTS if e['id'] == 'd1_nailart')
    if 'mehandi' in clean or 'mehndi' in clean: return next(e for e in EVENTS if e['id'] == 'd1_mehandi')
    if 'facepainting' in clean or 'facepaint' in clean or 'face' in clean: return next(e for e in EVENTS if e['id'] == 'd1_facepainting')
    if 'hackathon' in clean or 'intelligent' in clean or 'sustainability' in clean or 'fintech' in clean or 'smartbusiness' in clean or 'everyday' in clean: return next(e for e in EVENTS if e['id'] == 'd1_hackathon')
    if 'freefire' in clean or 'free fire' in clean or 'ff' in clean: return next(e for e in EVENTS if e['id'] == 'd1_freefire')
    if 'fixbug' in clean or 'bug' in clean or 'debugging' in clean: return next(e for e in EVENTS if e['id'] == 'd2_fixbug')
    if 'webforge' in clean or 'webdesign' in clean or 'webdev' in clean or 'web' in clean: return next(e for e in EVENTS if e['id'] == 'd2_webforge')
    if 'paperspark' in clean or 'paperpresentation' in clean or 'paper' in clean: return next(e for e in EVENTS if e['id'] == 'd2_paperspark')
    if 'prompting' in clean or 'promptwars' in clean or 'prompt' in clean: return next(e for e in EVENTS if e['id'] == 'd2_prompting')
    if 'connection' in clean or 'connections' in clean: return next(e for e in EVENTS if e['id'] == 'd2_connection')
    if 'solodance' in clean or ('dance' in clean and 'solo' in clean): return next(e for e in EVENTS if e['id'] == 'd2_solodance')
    if 'groupdance' in clean or ('dance' in clean and ('group' in clean or 'team' in clean)): return next(e for e in EVENTS if e['id'] == 'd2_groupdance')
    if 'singing' in clean or 'sing' in clean or 'solosinging' in clean: return next(e for e in EVENTS if e['id'] == 'd2_solosinging')
    if 'photo' in clean or 'photography' in clean: return next(e for e in EVENTS if e['id'] == 'd2_photography')
    if 'shortfilm' in clean or 'film' in clean: return next(e for e in EVENTS if e['id'] == 'd2_shortfilm')
    if 'reel' in clean or 'reels' in clean: return next(e for e in EVENTS if e['id'] == 'd2_reels')

    return None

def match_hackathon_theme(raw_str):
    if not raw_str: return 'Intelligent Systems'
    s = str(raw_str).lower()
    if 'intelligent' in s or 'ai' in s or 'smart system' in s: return 'Intelligent Systems'
    if 'sustainab' in s or 'green' in s or 'environment' in s: return 'Sustainability'
    if 'everyday' in s or 'daily' in s or 'iot' in s: return 'Tech for Everyday'
    if 'fintech' in s or 'commerce' in s or 'banking' in s or 'finance' in s: return 'FinTech / CommerceTech'
    if 'business' in s or 'institution' in s or 'automation' in s or 'mgmt' in s: return 'Smart Business Management / Institution Automation'
    return 'Intelligent Systems'

def extract_events_from_row(row_dict, default_event=None):
    matched = []
    seen = set()

    for k, v in row_dict.items():
        lk = str(k).lower()
        if 'event' in lk or 'competition' in lk or 'workshop' in lk:
            parts = re.split(r'[;,|\/\n]+', str(v))
            for p in parts:
                p_clean = p.strip()
                if p_clean:
                    m = match_event(p_clean)
                    if m and m['id'] not in seen:
                        seen.add(m['id'])
                        matched.append(m)

    if not matched and default_event:
        matched.append(default_event)
    return matched

@app.route('/api/admin/import/excel', methods=['POST'])
def import_excel():
    try:
        file_bytes = None
        if 'file' in request.files and request.files['file'].filename:
            file_bytes = request.files['file'].read()
        else:
            data = request.get_json(silent=True) or {}
            b64 = data.get('fileData') or data.get('fileBase64')
            if b64:
                import base64
                if ',' in b64:
                    b64 = b64.split(',')[1]
                file_bytes = base64.b64decode(b64)

        if not file_bytes:
            return jsonify({'success': False, 'message': 'No Excel file uploaded or file is empty.'}), 400

        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
        if not wb.sheetnames:
            return jsonify({'success': False, 'message': 'Excel workbook contains no sheets.'}), 400

        existing_codes_rows = query_all("SELECT reg_code FROM students UNION SELECT team_code FROM freefire_teams")
        existing_codes = set(str(r['reg_code']).strip().upper() for r in existing_codes_rows if r.get('reg_code'))

        idx_counter = [1]
        def gen_code(prefix='ONLINE-2026'):
            while True:
                code = f"{prefix}-{str(idx_counter[0]).zfill(4)}"
                idx_counter[0] += 1
                if code not in existing_codes:
                    existing_codes.add(code)
                    return code

        ff_idx_counter = [1]
        def gen_ff_code():
            while True:
                code = f"FF-TEAM-2026-{str(ff_idx_counter[0]).zfill(4)}"
                ff_idx_counter[0] += 1
                if code not in existing_codes:
                    existing_codes.add(code)
                    return code

        event_breakdown_map = {e['id']: {'eventId': e['id'], 'name': e['name'], 'day': e['day'].upper(), 'category': e['category'], 'count': 0} for e in EVENTS}

        total_students = 0
        total_online = 0
        total_spot = 0
        total_d1 = 0
        total_d2 = 0
        total_both = 0
        total_revenue = 0
        total_ff_players = 0
        total_ff_teams = 0
        sheets_processed = []

        conn = get_connection()
        try:
            conn.begin()
            with conn.cursor() as cur:
                for sheet_name in wb.sheetnames:
                    ws = wb[sheet_name]
                    rows_iter = ws.iter_rows(values_only=True)
                    try:
                        headers = next(rows_iter)
                    except StopIteration:
                        continue
                    if not headers:
                        continue

                    header_keys = [str(h).strip() if h is not None else f"col_{i}" for i, h in enumerate(headers)]
                    sheet_rows = []
                    for r in rows_iter:
                        if any(v is not None and str(v).strip() for v in r):
                            row_dict = {header_keys[i]: r[i] for i in range(min(len(header_keys), len(r)))}
                            sheet_rows.append(row_dict)

                    if not sheet_rows:
                        continue

                    sheets_processed.append(sheet_name)
                    is_ff_sheet = bool(re.search(r'free[\s_-]?fire', sheet_name, re.IGNORECASE))

                    if is_ff_sheet:
                        team_groups = {}
                        solo_players = []

                        for r in sheet_rows:
                            def get_v(keys):
                                for k, val in r.items():
                                    lk = re.sub(r'[^a-z0-9]', '', str(k).lower())
                                    for target in keys:
                                        if target in lk:
                                            return str(val).strip() if val is not None else ''
                                return ''

                            p_name = get_v(['name', 'player', 'participant'])
                            if not p_name: continue
                            p_phone = get_v(['phone', 'mobile', 'contact']) or '9999999999'
                            p_email = get_v(['email', 'mail']) or ''
                            p_college = get_v(['college', 'institute', 'university']) or 'Registered College'
                            p_team = get_v(['teamname', 'team', 'squad'])
                            p_cap = bool(re.search(r'yes|true|1|captain', get_v(['captain', 'iscaptain', 'role']), re.IGNORECASE))
                            p_source = 'SPOT' if re.search(r'spot', get_v(['type', 'source']), re.IGNORECASE) else 'ONLINE'

                            player_obj = {'name': p_name, 'phone': p_phone, 'email': p_email, 'college': p_college, 'is_captain': p_cap, 'registration_type': p_source}
                            if p_team:
                                team_groups.setdefault(p_team, []).append(player_obj)
                            else:
                                solo_players.append(player_obj)

                        for t_name, players in team_groups.items():
                            team_code = gen_ff_code()
                            source = players[0]['registration_type'] if players else 'ONLINE'
                            cur.execute("INSERT INTO freefire_teams (team_code, team_name, registration_source) VALUES (%s, %s, %s)", (team_code, t_name, source))
                            t_id = cur.lastrowid
                            total_ff_teams += 1

                            has_cap = any(p['is_captain'] for p in players)
                            for i, p in enumerate(players):
                                is_c = 1 if (p['is_captain'] or (not has_cap and i == 0)) else 0
                                cur.execute("INSERT INTO freefire_players (team_id, name, phone, email, college, registration_type, is_captain) VALUES (%s, %s, %s, %s, %s, %s, %s)",
                                            (t_id, p['name'], p['phone'], p['email'], p['college'], p['registration_type'], is_c))
                                total_ff_players += 1

                        for sp in solo_players:
                            cur.execute("INSERT INTO freefire_players (team_id, name, phone, email, college, registration_type, is_captain) VALUES (%s, %s, %s, %s, %s, %s, %s)",
                                        (None, sp['name'], sp['phone'], sp['email'], sp['college'], sp['registration_type'], 0))
                            total_ff_players += 1

                        continue

                    default_evt = match_event(sheet_name)

                    for r in sheet_rows:
                        def get_val(keys):
                            for k, val in r.items():
                                lk = re.sub(r'[^a-z0-9]', '', str(k).lower())
                                for target in keys:
                                    if target == lk or target in lk:
                                        return str(val).strip() if val is not None else ''
                            return ''

                        name = get_val(['name', 'studentname', 'participantname', 'fullname', 'candidate'])
                        if not name: continue

                        phone = get_val(['phone', 'phonenumber', 'mobile', 'mobilenumber', 'contact']) or '9999999999'
                        email = get_val(['email', 'emailaddress', 'emailid', 'mail']) or ''
                        college = get_val(['college', 'collegename', 'institute', 'institution', 'university']) or 'Registered College'
                        raw_code = get_val(['code', 'regcode', 'registrationcode', 'regid', 'token'])
                        raw_pass = get_val(['pass', 'passtype', 'day', 'dayselection'])
                        raw_fee = get_val(['fee', 'totalfee', 'feepaid', 'amount', 'paid'])
                        raw_source = get_val(['type', 'registrationtype', 'source'])
                        raw_theme = get_val(['theme', 'hackathontheme', 'topic', 'track'])
                        raw_food = get_val(['food', 'foodtoken'])
                        raw_tag = get_val(['tag', 'idtag', 'taggiven'])

                        matched_events = extract_events_from_row(r, default_evt)
                        if not matched_events:
                            matched_events = [EVENTS[0]]

                        has_ff = any(e['id'] == 'd1_freefire' for e in matched_events)
                        matched_events = [e for e in matched_events if e['id'] != 'd1_freefire']

                        if not matched_events and has_ff:
                            cur.execute("INSERT INTO freefire_players (team_id, name, phone, email, college, registration_type, is_captain) VALUES (%s, %s, %s, %s, %s, %s, %s)",
                                        (None, name, phone, email, college, 'SPOT' if re.search(r'spot', raw_source, re.IGNORECASE) else 'ONLINE', 0))
                            total_ff_players += 1
                            continue

                        if not matched_events:
                            matched_events = [EVENTS[0]]

                        pass_lower = raw_pass.lower()
                        if 'both' in pass_lower:
                            day_sel = 'both'
                        elif '2' in pass_lower or pass_lower == 'day2':
                            day_sel = 'day2'
                        elif '1' in pass_lower or pass_lower == 'day1':
                            day_sel = 'day1'
                        else:
                            has_d1 = any(e['day'] == 'day1' for e in matched_events)
                            has_d2 = any(e['day'] == 'day2' for e in matched_events)
                            if has_d1 and has_d2: day_sel = 'both'
                            elif has_d2: day_sel = 'day2'
                            else: day_sel = 'day1'

                        try:
                            fee = int(float(raw_fee))
                            if fee <= 0: fee = PRICING[day_sel]
                        except:
                            fee = PRICING[day_sel]

                        reg_type = 'SPOT' if re.search(r'spot', raw_source, re.IGNORECASE) else 'ONLINE'
                        food_val = 1 if re.search(r'yes|true|1', raw_food, re.IGNORECASE) else 0
                        tag_val = 1 if re.search(r'yes|true|1', raw_tag, re.IGNORECASE) else 0

                        final_reg_code = raw_code.upper() if raw_code else ''
                        if not final_reg_code or final_reg_code in existing_codes:
                            final_reg_code = gen_code('SPOT-2026' if reg_type == 'SPOT' else 'ONLINE-2026')
                        else:
                            existing_codes.add(final_reg_code)

                        theme_val = match_hackathon_theme(raw_theme) if any(e['id'] == 'd1_hackathon' for e in matched_events) else None

                        cur.execute(
                            "INSERT INTO students (reg_code, name, phone, college, day_selection, total_fee, email, registration_type, food_given, tag_given, food_d2_given, tag_d2_given) VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)",
                            (final_reg_code, name, phone, college, day_sel, fee, email, reg_type, food_val, tag_val, food_val, tag_val)
                        )
                        s_id = cur.lastrowid
                        total_students += 1
                        total_revenue += fee

                        if reg_type == 'SPOT': total_spot += 1
                        else: total_online += 1

                        if day_sel == 'day1': total_d1 += 1
                        elif day_sel == 'day2': total_d2 += 1
                        else: total_both += 1

                        for evt in matched_events:
                            cur.execute(
                                "INSERT INTO student_events (student_id, event_id, event_day, verified_status, hackathon_theme) VALUES (%s, %s, %s, %s, %s)",
                                (s_id, evt['id'], evt['day'], 0, theme_val if evt['id'] == 'd1_hackathon' else None)
                            )
                            if evt['id'] in event_breakdown_map:
                                event_breakdown_map[evt['id']]['count'] += 1

            conn.commit()

            active_breakdown = [eb for eb in event_breakdown_map.values() if eb['count'] > 0]
            return jsonify({
                'success': True,
                'message': f"Successfully imported {total_students} participants and split them across events!",
                'summary': {
                    'totalStudents': total_students,
                    'onlineCount': total_online,
                    'spotCount': total_spot,
                    'day1Count': total_d1,
                    'day2Count': total_d2,
                    'bothDaysCount': total_both,
                    'totalRevenue': total_revenue,
                    'freefirePlayers': total_ff_players,
                    'freefireTeams': total_ff_teams,
                    'eventBreakdown': active_breakdown,
                    'sheetsProcessed': sheets_processed
                }
            })
        except Exception as tx_err:
            conn.rollback()
            raise tx_err
        finally:
            conn.close()

    except Exception as e:
        print(f"Error importing Excel file: {e}")
        return jsonify({'success': False, 'message': f"Server error processing Excel sheet: {e}"}), 500

@app.route('/api/admin/import/sample-template', methods=['GET'])
def download_sample_template():
    try:
        wb = openpyxl.Workbook()
        ws_sample = wb.active
        ws_sample.title = "Participants_Import"

        sample_headers = ['Registration Code', 'Participant Name', 'Phone Number', 'Email Address', 'College / Institute', 'Pass Type', 'Registered Events', 'Hackathon Theme', 'Registration Type', 'Fee Paid (INR)']
        sample_rows = [
            ['ONLINE-202601', 'Kavya Sundaram', '9876543210', 'kavya.s@example.com', 'PSG College of Technology', 'BOTH', 'Adzap; Webforge', '', 'ONLINE', 350],
            ['ONLINE-202602', 'Arun Kumar M', '9840123456', 'arun.k@example.com', 'CIT Coimbatore', 'DAY1', 'Hackathon; Quiz', 'Intelligent Systems', 'ONLINE', 250],
            ['ONLINE-202603', 'Pooja R', '9790234567', 'pooja.r@example.com', 'Anna University', 'DAY2', 'Solo Dance; Photography Competition', '', 'ONLINE', 250]
        ]

        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="4F46E5", end_color="4F46E5", fill_type="solid")

        ws_sample.append(sample_headers)
        for cell in ws_sample[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for row in sample_rows:
            ws_sample.append(row)
        for col in ws_sample.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws_sample.column_dimensions[col_letter].width = max(max_len + 3, 14)

        ws_ref = wb.create_sheet("Event_Catalog_Reference")
        ref_headers = ['#', 'Event Name', 'Event ID', 'Day', 'Category', 'Min Participants', 'Max Participants', 'Notes']
        ref_rows = [[i + 1, e['name'], e['id'], e['day'].upper(), e['category'], e.get('min_participants', 1), e.get('max_participants', 1), 'Themes: Intelligent Systems, Sustainability, Tech for Everyday, FinTech / CommerceTech, Smart Business Management' if e['id'] == 'd1_hackathon' else ('Standalone Tournament' if e.get('is_standalone') else 'Fest Event')] for i, e in enumerate(EVENTS)]

        ws_ref.append(ref_headers)
        for cell in ws_ref[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for row in ref_rows:
            ws_ref.append(row)
        for col in ws_ref.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws_ref.column_dimensions[col_letter].width = max(max_len + 3, 14)

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return send_file(buf, as_attachment=True, download_name="TechRaga26_Import_Template.xlsx", mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    except Exception as e:
        print(f"Error generating template: {e}")
        return jsonify({'success': False, 'message': 'Failed to generate template.'}), 500


# ==========================================
# 10. FREE FIRE ESPORTS TOURNAMENT
# ==========================================
@app.route('/api/freefire/spot-register', methods=['POST'])
def freefire_spot_register():
    try:
        data = request.get_json() or {}
        team_name = (data.get('teamName') or '').strip()
        members = data.get('members') or data.get('players') or []
        counter_name = data.get('counterName')
        faculty_name = data.get('facultyName')

        if not team_name:
            return jsonify({'success': False, 'message': 'Team Name is required.'}), 400
        if not isinstance(members, list) or len(members) != 4:
            return jsonify({'success': False, 'message': 'Exactly 4 Team Members are required for Free Fire.'}), 400

        for i, m in enumerate(members):
            if not m.get('name') or not str(m['name']).strip():
                return jsonify({'success': False, 'message': f"Player {i + 1} Name is required."}), 400
            if not m.get('phone') or not re.match(r'^\d{10}$', str(m['phone']).strip()):
                return jsonify({'success': False, 'message': f"Player {i + 1} valid 10-digit Phone is required."}), 400
            if not m.get('email') or not str(m['email']).strip():
                return jsonify({'success': False, 'message': f"Player {i + 1} Email Address is required."}), 400
            if not m.get('college') or not str(m['college']).strip():
                return jsonify({'success': False, 'message': f"Player {i + 1} College is required."}), 400

        final_counter = str(counter_name).strip() if counter_name and str(counter_name).strip() else None
        final_faculty = str(faculty_name).strip() if faculty_name and str(faculty_name).strip() else None
        if final_counter and not final_faculty:
            c_row = query_get("SELECT faculty_name FROM counters WHERE counter_name = ?", (final_counter,))
            if c_row and c_row.get('faculty_name'):
                final_faculty = c_row['faculty_name']

        random_num = random.randint(1000, 9999)
        team_code = f"FF-TEAM-2026-{random_num}"
        reg_code = f"SPOT-2026-{random_num}"

        conn = get_connection()
        try:
            conn.begin()
            with conn.cursor() as cur:
                cur.execute(
                    "INSERT INTO freefire_teams (team_code, team_name, registration_source, counter_name, faculty_name) VALUES (%s, %s, %s, %s, %s)",
                    (team_code, team_name, 'SPOT', final_counter, final_faculty)
                )
                team_id = cur.lastrowid

                for i, m in enumerate(members):
                    cur.execute(
                        "INSERT INTO freefire_players (team_id, name, phone, email, college, registration_type, is_captain, counter_name, faculty_name) VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s)",
                        (team_id, str(m['name']).strip(), str(m['phone']).strip(), str(m.get('email', '')).strip(), str(m['college']).strip(), 'SPOT', 1 if i == 0 else 0, final_counter, final_faculty)
                    )

                cur.execute(
                    "INSERT INTO students (reg_code, name, phone, email, college, day_selection, total_fee, registration_type, counter_name, faculty_name) VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s)",
                    (reg_code, f"{team_name} (Captain: {str(members[0]['name']).strip()})", str(members[0]['phone']).strip(), str(members[0].get('email', '')).strip(), str(members[0]['college']).strip(), 'day1', 400, 'SPOT', final_counter, final_faculty)
                )
                student_id = cur.lastrowid

                cur.execute(
                    "INSERT INTO student_events (student_id, event_id, event_day, verified_status, hackathon_theme) VALUES (%s, %s, %s, %s, %s)",
                    (student_id, 'd1_freefire', 'day1', 0, None)
                )

            conn.commit()
            return jsonify({
                'success': True,
                'message': 'Free Fire Team Spot Registration completed successfully!',
                'teamId': team_id,
                'teamCode': team_code,
                'team': {
                    'id': team_id,
                    'team_id': team_id,
                    'team_code': team_code,
                    'teamCode': team_code,
                    'team_name': team_name,
                    'registration_source': 'SPOT',
                    'counter_name': final_counter,
                    'faculty_name': final_faculty,
                    'members': members
                }
            })
        except Exception as tx_err:
            conn.rollback()
            raise tx_err
        finally:
            conn.close()

    except Exception as e:
        print(f"Free Fire registration error: {e}")
        return jsonify({'success': False, 'message': 'Failed to process Free Fire team registration.'}), 500

@app.route('/api/freefire/unassigned-online-players', methods=['GET'])
def get_unassigned_freefire_players():
    try:
        players = query_all("SELECT id, name, phone, email, college, registration_type, registered_at FROM freefire_players WHERE team_id IS NULL ORDER BY registered_at ASC")
        return jsonify({'success': True, 'count': len(players), 'players': players})
    except Exception as e:
        print(f"Error fetching online players: {e}")
        return jsonify({'success': False, 'message': 'Failed to fetch online solo players.'}), 500

@app.route('/api/freefire/create-team', methods=['POST'])
def create_freefire_team():
    try:
        data = request.get_json() or {}
        team_name = (data.get('teamName') or '').strip()
        player_ids = data.get('playerIds', [])

        if not team_name:
            return jsonify({'success': False, 'message': 'Team Name is required.'}), 400
        if not isinstance(player_ids, list) or len(player_ids) != 4:
            return jsonify({'success': False, 'message': 'Exactly 4 player IDs must be selected to form a team.'}), 400

        random_num = random.randint(1000, 9999)
        team_code = f"FF-TEAM-2026-{random_num}"

        conn = get_connection()
        try:
            conn.begin()
            with conn.cursor() as cur:
                cur.execute("INSERT INTO freefire_teams (team_code, team_name, registration_source) VALUES (%s, %s, %s)", (team_code, team_name, 'COORDINATOR_FORMED'))
                team_id = cur.lastrowid

                for i, pid in enumerate(player_ids):
                    cur.execute("UPDATE freefire_players SET team_id = %s, is_captain = %s WHERE id = %s", (team_id, 1 if i == 0 else 0, pid))

            conn.commit()
            return jsonify({'success': True, 'message': 'Free Fire Team formed successfully!'})
        except Exception as tx_err:
            conn.rollback()
            raise tx_err
        finally:
            conn.close()

    except Exception as e:
        print(f"Error forming team: {e}")
        return jsonify({'success': False, 'message': 'Failed to create team.'}), 500

@app.route('/api/freefire/teams', methods=['GET'])
def get_freefire_teams():
    try:
        rows = query_all("""
            SELECT 
              t.id as team_id, t.team_code, t.team_name, t.registration_source, t.registered_at as team_registered_at,
              p.id as player_id, p.name as player_name, p.phone as player_phone, p.email as player_email, p.college as player_college, p.is_captain
            FROM freefire_teams t
            JOIN freefire_players p ON t.id = p.team_id
            ORDER BY t.registered_at DESC, t.id DESC, p.is_captain DESC
        """)

        team_map = {}
        for r in rows:
            tid = r['team_id']
            if tid not in team_map:
                team_map[tid] = {
                    'id': r['team_id'],
                    'team_code': r['team_code'],
                    'team_name': r['team_name'],
                    'registration_source': r['registration_source'],
                    'registered_at': r['team_registered_at'],
                    'members': []
                }
            team_map[tid]['members'].append({
                'id': r['player_id'],
                'name': r['player_name'],
                'phone': r['player_phone'],
                'email': r.get('player_email') or '',
                'college': r['player_college'],
                'is_captain': bool(r['is_captain'])
            })

        teams = list(team_map.values())
        return jsonify({'success': True, 'count': len(teams), 'teams': teams})
    except Exception as e:
        print(f"Error fetching Free Fire teams: {e}")
        return jsonify({'success': False, 'message': 'Failed to fetch teams.'}), 500

@app.route('/api/freefire/teams/<int:team_id>', methods=['PUT'])
def edit_freefire_team(team_id):
    try:
        data = request.get_json() or {}
        team_name = (data.get('teamName') or '').strip()
        members = data.get('members', [])
        removed_member_ids = data.get('removedMemberIds', [])
        add_player_ids = data.get('addPlayerIds', [])
        new_members = data.get('newMembers', [])

        if not team_name:
            return jsonify({'success': False, 'message': 'Team Name is required.'}), 400

        conn = get_connection()
        try:
            conn.begin()
            with conn.cursor() as cur:
                cur.execute("UPDATE freefire_teams SET team_name = %s WHERE id = %s", (team_name, team_id))

                if isinstance(removed_member_ids, list):
                    for rem_id in removed_member_ids:
                        cur.execute("UPDATE freefire_players SET team_id = NULL, is_captain = 0 WHERE id = %s AND team_id = %s", (rem_id, team_id))

                if isinstance(members, list):
                    for i, m in enumerate(members):
                        if m.get('id'):
                            cur.execute(
                                "UPDATE freefire_players SET name = %s, phone = %s, email = %s, college = %s, is_captain = %s WHERE id = %s AND team_id = %s",
                                (str(m['name']).strip(), str(m.get('phone', '')).strip(), str(m.get('email', '')).strip(), str(m.get('college', '')).strip(), 1 if i == 0 else 0, m['id'], team_id)
                            )

                if isinstance(add_player_ids, list):
                    for pid in add_player_ids:
                        cur.execute("UPDATE freefire_players SET team_id = %s, is_captain = 0 WHERE id = %s", (team_id, pid))

                if isinstance(new_members, list):
                    for nm in new_members:
                        if nm.get('name') and str(nm['name']).strip():
                            cur.execute(
                                "INSERT INTO freefire_players (team_id, name, phone, email, college, registration_type, is_captain) VALUES (%s, %s, %s, %s, %s, %s, %s)",
                                (team_id, str(nm['name']).strip(), str(nm.get('phone', '')).strip(), str(nm.get('email', '')).strip(), str(nm.get('college', '')).strip(), 'SPOT', 0)
                            )

                cur.execute("SELECT id, is_captain FROM freefire_players WHERE team_id = %s ORDER BY id ASC", (team_id,))
                team_players = cur.fetchall()
                if team_players and not any(p['is_captain'] == 1 for p in team_players):
                    cur.execute("UPDATE freefire_players SET is_captain = 1 WHERE id = %s", (team_players[0]['id'],))

            conn.commit()
            return jsonify({'success': True, 'message': 'Free Fire Team roster updated successfully!'})
        except Exception as tx_err:
            conn.rollback()
            raise tx_err
        finally:
            conn.close()

    except Exception as e:
        print(f"Error editing team: {e}")
        return jsonify({'success': False, 'message': 'Failed to update team details.'}), 500


# ==========================================
# 11. COUNTERS & FACULTY MANAGEMENT
# ==========================================
@app.route('/api/counters', methods=['GET'])
def get_counters():
    try:
        counters = query_all("SELECT id, counter_name, faculty_name, status, updated_at FROM counters ORDER BY id ASC")
        stats_rows = query_all("""
            SELECT counter_name, COUNT(*) as total_registrations, SUM(total_fee) as total_amount
            FROM students
            WHERE counter_name IS NOT NULL AND counter_name != ''
            GROUP BY counter_name
        """)
        stats_map = {r['counter_name']: {'total_registrations': r['total_registrations'] or 0, 'total_amount': r['total_amount'] or 0} for r in stats_rows}

        enriched = []
        for c in counters:
            st = stats_map.get(c['counter_name'], {'total_registrations': 0, 'total_amount': 0})
            c_dict = dict(c)
            c_dict['faculty_name'] = c.get('faculty_name') or ''
            c_dict['status'] = c.get('status') or 'ACTIVE'
            c_dict['total_registrations'] = st['total_registrations']
            c_dict['total_amount'] = st['total_amount']
            enriched.append(c_dict)

        return jsonify({'success': True, 'counters': enriched})
    except Exception as e:
        print(f"Error fetching counters: {e}")
        return jsonify({'success': False, 'message': 'Failed to fetch counters list.'}), 500

@app.route('/api/faculty-list', methods=['GET'])
def get_faculty_list():
    try:
        faculty_rows = query_all("SELECT id, name, department, created_at FROM faculty_members ORDER BY name ASC")
        counters_rows = query_all("SELECT counter_name, faculty_name FROM counters WHERE faculty_name IS NOT NULL AND faculty_name != ''")
        assigned_map = {c['faculty_name']: c['counter_name'] for c in counters_rows}

        faculty = [{
            'id': f['id'],
            'name': f['name'],
            'department': f.get('department') or 'Academics',
            'assigned_counter': assigned_map.get(f['name'])
        } for f in faculty_rows]

        return jsonify({'success': True, 'faculty': faculty})
    except Exception as e:
        print(f"Error fetching faculty list: {e}")
        return jsonify({'success': False, 'message': 'Failed to fetch faculty list.'}), 500

@app.route('/api/counters/assign', methods=['POST'])
def assign_faculty_to_counter():
    try:
        data = request.get_json() or {}
        counter_id = data.get('counterId')
        counter_name = data.get('counterName')
        faculty_name = data.get('facultyName')

        target_counter = None
        if counter_id:
            target_counter = query_get("SELECT id, counter_name, faculty_name FROM counters WHERE id = ?", (counter_id,))
        elif counter_name:
            target_counter = query_get("SELECT id, counter_name, faculty_name FROM counters WHERE counter_name = ?", (str(counter_name).strip(),))

        if not target_counter:
            return jsonify({'success': False, 'message': 'Registration Counter not found.'}), 404

        new_faculty = str(faculty_name).strip() if faculty_name and str(faculty_name).strip() else None

        if new_faculty:
            existing = query_get("SELECT id, counter_name, faculty_name FROM counters WHERE faculty_name = ? AND id != ?", (new_faculty, target_counter['id']))
            if existing:
                return jsonify({
                    'success': False,
                    'message': f"Faculty '{new_faculty}' is already assigned to {existing['counter_name']}. One faculty member can only be assigned to one counter."
                }), 400

            query_run("INSERT IGNORE INTO faculty_members (name) VALUES (?)", (new_faculty,))

        query_run("UPDATE counters SET faculty_name = ?, updated_at = CURRENT_TIMESTAMP WHERE id = ?", (new_faculty, target_counter['id']))

        return jsonify({
            'success': True,
            'message': f"{new_faculty} successfully assigned to {target_counter['counter_name']}!" if new_faculty else f"Faculty unassigned from {target_counter['counter_name']}.",
            'counter': {
                'id': target_counter['id'],
                'counter_name': target_counter['counter_name'],
                'faculty_name': new_faculty or ''
            }
        })
    except Exception as e:
        print(f"Error assigning faculty: {e}")
        return jsonify({'success': False, 'message': 'Server error while assigning faculty to counter.'}), 500

@app.route('/api/faculty/add', methods=['POST'])
def add_faculty_member():
    try:
        data = request.get_json() or {}
        name = (data.get('name') or '').strip()
        dept = (data.get('department') or 'Academics').strip()

        if not name:
            return jsonify({'success': False, 'message': 'Faculty Name is required.'}), 400

        existing = query_get("SELECT id FROM faculty_members WHERE name = ?", (name,))
        if existing:
            return jsonify({'success': False, 'message': f"Faculty '{name}' already exists in catalog."}), 400

        res = query_run("INSERT INTO faculty_members (name, department) VALUES (?, ?)", (name, dept))
        return jsonify({
            'success': True,
            'message': f"Faculty '{name}' added successfully.",
            'faculty': {'id': res['lastID'], 'name': name, 'department': dept}
        })
    except Exception as e:
        print(f"Error adding faculty: {e}")
        return jsonify({'success': False, 'message': 'Failed to add faculty member.'}), 500

@app.route('/api/admin/counter-report', methods=['GET'])
def get_counter_report():
    try:
        counter_filter = request.args.get('counter')
        faculty_filter = request.args.get('faculty')
        date_filter = request.args.get('date')
        q = request.args.get('q', '').strip()

        all_counters = query_all("SELECT id, counter_name, faculty_name, status FROM counters ORDER BY id ASC")

        sql = """
          SELECT 
            s.id, s.reg_code, s.name, s.phone, s.email, s.college, s.day_selection, s.total_fee, s.registration_type,
            s.counter_name, s.faculty_name, s.food_given, s.tag_given, s.registered_at,
            se.event_id, se.event_day, se.verified_status, se.hackathon_theme,
            e.name as event_name, e.category as event_category
          FROM students s
          LEFT JOIN student_events se ON s.id = se.student_id
          LEFT JOIN events e ON se.event_id = e.id
          WHERE UPPER(COALESCE(s.registration_type, 'ONLINE')) = 'SPOT'
        """
        params = []

        if counter_filter and counter_filter != 'all':
            sql += " AND s.counter_name = ?"
            params.append(counter_filter.strip())

        if faculty_filter and faculty_filter != 'all':
            sql += " AND s.faculty_name = ?"
            params.append(faculty_filter.strip())

        if date_filter and date_filter.strip():
            sql += " AND DATE(s.registered_at) = DATE(?)"
            params.append(date_filter.strip())

        if q:
            term = f"%{q}%"
            sql += " AND (s.name LIKE ? OR s.phone LIKE ? OR s.college LIKE ? OR s.reg_code LIKE ? OR s.email LIKE ?)"
            params.extend([term, term, term, term, term])

        sql += " ORDER BY s.registered_at DESC, s.id DESC"

        rows = query_all(sql, params)

        student_map = {}
        for r in rows:
            sid = r['id']
            if sid not in student_map:
                student_map[sid] = {
                    'id': r['id'],
                    'reg_code': r['reg_code'],
                    'name': r['name'],
                    'phone': r['phone'],
                    'email': r.get('email') or '',
                    'college': r['college'],
                    'day_selection': r['day_selection'],
                    'total_fee': r['total_fee'],
                    'counter_name': r.get('counter_name') or 'Unassigned Counter',
                    'faculty_name': r.get('faculty_name') or 'Unassigned Faculty',
                    'food_given': bool(r['food_given']),
                    'tag_given': bool(r['tag_given']),
                    'registered_at': r['registered_at'],
                    'events': []
                }
            if r.get('event_name'):
                student_map[sid]['events'].append({
                    'event_id': r['event_id'],
                    'event_name': r['event_name'],
                    'event_day': r['event_day'],
                    'hackathon_theme': r.get('hackathon_theme')
                })

        filtered_students = list(student_map.values())

        counter_stats_rows = query_all("""
            SELECT 
              COALESCE(counter_name, 'Unassigned') as counter_name,
              COALESCE(faculty_name, 'Unassigned') as faculty_name,
              COUNT(*) as total_registrations,
              SUM(total_fee) as total_amount
            FROM students
            WHERE UPPER(COALESCE(registration_type, 'ONLINE')) = 'SPOT'
            GROUP BY counter_name
        """)
        counter_stats_map = {cs['counter_name']: {'total_registrations': cs['total_registrations'] or 0, 'total_amount': cs['total_amount'] or 0, 'faculty_name': cs['faculty_name']} for cs in counter_stats_rows}

        counter_wise_counts = []
        for c in all_counters:
            st = counter_stats_map.get(c['counter_name'], {'total_registrations': 0, 'total_amount': 0, 'faculty_name': 'Not Assigned'})
            counter_wise_counts.append({
                'id': c['id'],
                'counter_name': c['counter_name'],
                'faculty_name': c.get('faculty_name') or st.get('faculty_name') or 'Not Assigned',
                'total_registrations': st['total_registrations'],
                'total_amount': st['total_amount'],
                'status': c.get('status') or 'ACTIVE'
            })

        overall_row = query_get("SELECT COUNT(*) as count, SUM(total_fee) as total FROM students WHERE UPPER(COALESCE(registration_type, 'ONLINE')) = 'SPOT'")
        tot_spot = overall_row['count'] if overall_row else 0
        tot_amount = overall_row['total'] if overall_row and overall_row['total'] is not None else 0

        total_counters_active = len([c for c in counter_wise_counts if (c['faculty_name'] and c['faculty_name'] != 'Not Assigned') or c['total_registrations'] > 0])

        highest_counter = None
        max_c = -1
        for c in counter_wise_counts:
            if c['total_registrations'] > max_c and c['total_registrations'] > 0:
                max_c = c['total_registrations']
                highest_counter = {
                    'counter_name': c['counter_name'],
                    'faculty_name': c['faculty_name'],
                    'total_registrations': c['total_registrations'],
                    'total_amount': c['total_amount']
                }

        if not highest_counter and counter_wise_counts:
            highest_counter = {
                'counter_name': counter_wise_counts[0]['counter_name'],
                'faculty_name': counter_wise_counts[0]['faculty_name'],
                'total_registrations': counter_wise_counts[0]['total_registrations'],
                'total_amount': counter_wise_counts[0]['total_amount']
            }

        return jsonify({
            'success': True,
            'summary': {
                'totalCountersActive': total_counters_active,
                'totalSpotRegistrations': tot_spot,
                'totalAmountCollected': tot_amount,
                'highestPerformingCounter': highest_counter,
                'filteredCount': len(filtered_students),
                'filteredAmount': sum(s['total_fee'] for s in filtered_students)
            },
            'counterWiseCounts': counter_wise_counts,
            'registrations': filtered_students
        })
    except Exception as e:
        print(f"Error generating counter report: {e}")
        return jsonify({'success': False, 'message': 'Failed to generate counter report.'}), 500

@app.route('/api/admin/counter-report/export/excel', methods=['GET'])
def export_counter_report_excel():
    try:
        all_counters = query_all("SELECT id, counter_name, faculty_name FROM counters ORDER BY id ASC")
        counter_stats_rows = query_all("SELECT COALESCE(counter_name, 'Unassigned') as counter_name, COUNT(*) as total_registrations, SUM(total_fee) as total_amount FROM students WHERE UPPER(COALESCE(registration_type, 'ONLINE')) = 'SPOT' GROUP BY counter_name")
        counter_stats_map = {cs['counter_name']: {'total_registrations': cs['total_registrations'] or 0, 'total_amount': cs['total_amount'] or 0} for cs in counter_stats_rows}

        summary_data = []
        for c in all_counters:
            st = counter_stats_map.get(c['counter_name'], {'total_registrations': 0, 'total_amount': 0})
            reg_cnt = st['total_registrations']
            amt = st['total_amount']
            summary_data.append([c['counter_name'], c.get('faculty_name') or 'Unassigned', reg_cnt, amt, round(amt / reg_cnt) if reg_cnt > 0 else 0])

        student_rows = query_all("""
            SELECT 
              s.id, s.reg_code, s.name, s.phone, s.email, s.college, s.day_selection, s.total_fee,
              s.counter_name, s.faculty_name, s.food_given, s.tag_given, s.registered_at,
              se.event_name, se.hackathon_theme
            FROM students s
            LEFT JOIN (
              SELECT se2.student_id, e2.name as event_name, se2.hackathon_theme
              FROM student_events se2
              JOIN events e2 ON se2.event_id = e2.id
            ) se ON s.id = se.student_id
            WHERE UPPER(COALESCE(s.registration_type, 'ONLINE')) = 'SPOT'
            ORDER BY s.counter_name ASC, s.registered_at DESC
        """)

        student_map = {}
        for r in student_rows:
            sid = r['id']
            if sid not in student_map:
                student_map[sid] = {
                    'reg_code': r['reg_code'], 'name': r['name'], 'phone': r['phone'], 'email': r.get('email') or 'N/A',
                    'college': r['college'], 'counter_name': r.get('counter_name') or 'Unassigned',
                    'faculty_name': r.get('faculty_name') or 'Unassigned',
                    'day_selection': (r.get('day_selection') or '').upper(),
                    'total_fee': r['total_fee'],
                    'food_given': 'YES' if r['food_given'] == 1 else 'NO',
                    'tag_given': 'YES' if r['tag_given'] == 1 else 'NO',
                    'registered_at': str(r['registered_at']),
                    'events': []
                }
            if r.get('event_name'):
                e_str = r['event_name']
                if r.get('hackathon_theme'): e_str += f" [{r['hackathon_theme']}]"
                student_map[sid]['events'].append(e_str)

        all_spot_rows = [[i + 1, s['reg_code'], s['name'], s['phone'], s['email'], s['college'], s['counter_name'], s['faculty_name'], s['day_selection'], s['total_fee'], s['food_given'], s['tag_given'], '; '.join(s['events']) or 'None', s['registered_at']] for i, s in enumerate(student_map.values())]

        wb = openpyxl.Workbook()
        ws_sum = wb.active
        ws_sum.title = "Counter_Summary"

        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="4F46E5", end_color="4F46E5", fill_type="solid")

        def style_and_add(ws, headers, rows):
            ws.append(headers)
            for cell in ws[1]:
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal="center", vertical="center")
            for r in rows:
                ws.append(r)
            for col in ws.columns:
                max_len = max(len(str(cell.value or '')) for cell in col)
                col_letter = get_column_letter(col[0].column)
                ws.column_dimensions[col_letter].width = max(max_len + 3, 14)

        style_and_add(ws_sum, ['Counter Number', 'Assigned Faculty Name', 'Total Spot Registrations', 'Total Amount Collected (INR)', 'Average Collection / Reg (INR)'], summary_data)

        ws_roster = wb.create_sheet("All_Spot_Registrations")
        style_and_add(ws_roster, ['S.No', 'Registration Code', 'Participant Name', 'Phone Number', 'Email Address', 'College / Institution', 'Counter Number', 'Faculty In-Charge', 'Pass Type', 'Registration Fee (INR)', 'Food Token', 'ID Tag', 'Registered Events', 'Registration Time'], all_spot_rows)

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return send_file(buf, as_attachment=True, download_name="TechRaga26_Counter_Registration_Report.xlsx", mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    except Exception as e:
        print(f"Error exporting counter excel: {e}")
        return jsonify({'success': False, 'message': 'Failed to export counter report to Excel.'}), 500

@app.route('/api/admin/counter-report/export/overall', methods=['GET'])
def export_overall_report():
    try:
        student_rows = query_all("""
            SELECT 
              s.id, s.reg_code, s.name, s.phone, s.email, s.college, s.day_selection, s.total_fee, s.registration_type,
              s.counter_name, s.faculty_name, s.food_given, s.tag_given, s.registered_at,
              se.event_name, se.hackathon_theme
            FROM students s
            LEFT JOIN (
              SELECT se2.student_id, e2.name as event_name, se2.hackathon_theme
              FROM student_events se2
              JOIN events e2 ON se2.event_id = e2.id
            ) se ON s.id = se.student_id
            ORDER BY s.id ASC
        """)

        student_map = {}
        for r in student_rows:
            sid = r['id']
            if sid not in student_map:
                student_map[sid] = {
                    'reg_code': r['reg_code'], 'name': r['name'], 'phone': r['phone'], 'email': r.get('email') or 'N/A',
                    'college': r['college'], 'source': (r.get('registration_type') or 'ONLINE').upper(),
                    'counter': r.get('counter_name') or ('Unassigned' if r.get('registration_type') == 'SPOT' else 'Online Portal'),
                    'faculty': r.get('faculty_name') or ('Unassigned' if r.get('registration_type') == 'SPOT' else 'N/A'),
                    'pass': (r.get('day_selection') or '').upper(),
                    'total_fee': r['total_fee'],
                    'registered_at': str(r['registered_at']),
                    'events': []
                }
            if r.get('event_name'):
                e_str = r['event_name']
                if r.get('hackathon_theme'): e_str += f" [{r['hackathon_theme']}]"
                student_map[sid]['events'].append(e_str)

        master_rows = [[i + 1, s['reg_code'], s['name'], s['phone'], s['email'], s['college'], s['source'], s['counter'], s['faculty'], s['pass'], s['total_fee'], '; '.join(s['events']) or 'None', s['registered_at']] for i, s in enumerate(student_map.values())]

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Overall_Registrations"

        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="4F46E5", end_color="4F46E5", fill_type="solid")

        ws.append(['S.No', 'Registration Code', 'Participant Name', 'Phone Number', 'Email Address', 'College', 'Source', 'Counter', 'Faculty In-Charge', 'Pass', 'Amount (INR)', 'Registered Events', 'Registration Date'])
        for cell in ws[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for r in master_rows:
            ws.append(r)
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 14)

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return send_file(buf, as_attachment=True, download_name="TechRaga26_Overall_Registration_Report.xlsx", mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    except Exception as e:
        print(f"Error generating overall report: {e}")
        return jsonify({'success': False, 'message': 'Failed to generate overall report.'}), 500


# ==========================================
# 12. ACCESS CONTROL & AUTHENTICATION
# ==========================================
@app.route('/api/auth/verify', methods=['POST'])
def verify_auth():
    try:
        data = request.get_json() or {}
        module_key = (data.get('moduleKey') or '').strip()
        username = (data.get('username') or '').strip()
        password = (data.get('password') or '').strip()

        if not username or not password:
            return jsonify({'success': False, 'message': 'Username and password are required.'}), 400

        u = username.lower()
        p = password

        # 1. Master Admin Check (module_key == 'admin_master' or login as admin)
        admin_entry = next((c for c in default_creds if get_cred_val(c, 'key', 'module_key') == 'admin_master'), None)
        if admin_entry and get_cred_val(admin_entry, 'user', 'username').lower() == u and get_cred_val(admin_entry, 'pass', 'password') == p:
            return jsonify({
                'success': True,
                'message': 'Master Admin Access Granted',
                'role': 'ADMIN',
                'moduleKey': module_key or 'admin_master',
                'moduleName': 'Admin Master Portal',
                'username': get_cred_val(admin_entry, 'user', 'username')
            })

        # 2. Specific Module Check
        target_cred = None
        if module_key:
            target_cred = next((c for c in default_creds if get_cred_val(c, 'key', 'module_key') == module_key and get_cred_val(c, 'user', 'username').lower() == u and get_cred_val(c, 'pass', 'password') == p), None)
        else:
            target_cred = next((c for c in default_creds if get_cred_val(c, 'user', 'username').lower() == u and get_cred_val(c, 'pass', 'password') == p), None)

        # 3. Fallback for coordinator portal
        if not target_cred and (module_key == 'coordinator_portal' or not module_key):
            target_cred = next((c for c in default_creds if get_cred_val(c, 'user', 'username').lower() == u and get_cred_val(c, 'pass', 'password') == p), None)

        if target_cred:
            cat = get_cred_val(target_cred, 'category', default='PORTAL')
            m_key = get_cred_val(target_cred, 'key', 'module_key')
            m_name = get_cred_val(target_cred, 'name', 'module_name')
            m_user = get_cred_val(target_cred, 'user', 'username')
            return jsonify({
                'success': True,
                'message': f"Welcome to {m_name}",
                'role': 'PORTAL_OPERATOR' if cat == 'PORTAL' else 'EVENT_COORDINATOR',
                'moduleKey': m_key,
                'moduleName': m_name,
                'category': cat,
                'username': m_user
            })

        return jsonify({'success': False, 'message': 'Invalid username or password.'}), 401
    except Exception as e:
        print(f"Authentication error: {e}")
        return jsonify({'success': False, 'message': 'Authentication verification failed.'}), 500

@app.route('/api/admin/credentials', methods=['GET'])
def get_admin_credentials():
    try:
        credentials = []
        for i, c in enumerate(default_creds, 1):
            credentials.append({
                'id': i,
                'module_key': get_cred_val(c, 'key', 'module_key'),
                'module_name': get_cred_val(c, 'name', 'module_name'),
                'category': get_cred_val(c, 'category', default='PORTAL'),
                'username': get_cred_val(c, 'user', 'username'),
                'password': get_cred_val(c, 'pass', 'password'),
                'updated_at': datetime.now().strftime('%Y-%m-%d %H:%M:%S')
            })
        return jsonify({'success': True, 'credentials': credentials})
    except Exception as e:
        print(f"Error fetching credentials: {e}")
        return jsonify({'success': False, 'message': 'Failed to fetch credentials list.'}), 500

@app.route('/api/admin/credentials/update', methods=['POST'])
def update_admin_credential():
    try:
        data = request.get_json() or {}
        module_key = data.get('moduleKey')
        username = (data.get('username') or '').strip()
        password = (data.get('password') or '').strip()

        if not module_key or not username or not password:
            return jsonify({'success': False, 'message': 'Module Key, Username, and Password are required.'}), 400

        entry = next((c for c in default_creds if get_cred_val(c, 'key', 'module_key') == module_key), None)
        if not entry:
            return jsonify({'success': False, 'message': f"No credential record found for key: {module_key}"}), 404

        if 'user' in entry: entry['user'] = username
        if 'username' in entry: entry['username'] = username
        if 'pass' in entry: entry['pass'] = password
        if 'password' in entry: entry['password'] = password

        updated = {
            'module_key': get_cred_val(entry, 'key', 'module_key'),
            'module_name': get_cred_val(entry, 'name', 'module_name'),
            'category': get_cred_val(entry, 'category', default='PORTAL'),
            'username': username,
            'password': password
        }

        return jsonify({
            'success': True,
            'message': f"Credentials for '{updated['module_name']}' updated successfully!",
            'credential': updated
        })
    except Exception as e:
        print(f"Error updating credentials: {e}")
        return jsonify({'success': False, 'message': 'Failed to update credentials.'}), 500


# ==========================================
# MAIN RUNNER
# ==========================================
if __name__ == '__main__':
    try:
        init_db_pool()
    except Exception as err:
        print(f"[WARNING] Database initialization warning: {err}")
    print(f"[READY] Python Flask Event Spot Registration Server running at http://localhost:{PORT}")
    app.run(host='0.0.0.0', port=PORT, debug=True)
